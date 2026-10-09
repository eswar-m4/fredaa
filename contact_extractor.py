#!/usr/bin/env python3
"""
Contact Extractor — single-file command-line version
======================================================

Excel in, Excel out. Every input URL is opened in headless Chromium via the
Python Playwright package (no Node). Links are scored before they are opened:
team, leadership, and staff pages outrank contact-us pages, which usually
publish an address and not people's names and titles. Once those URLs are
shortlisted and fetched, the page text goes straight to OpenAI. There is no
rule-based contact pass (schema.org, person cards, or mailto pairing) between
the shortlist and the model. Every field the model returns is checked against
the source text before being trusted.

This is the whole tool combined into one file for easy command-line use.
The original multi-file package also had some provisions for being called
from an embedding application later (a --config YAML file, --resume/
checkpoint support so a long batch survives an interruption, and an NDJSON
progress stream on stdout for a GUI to parse). Those aren't needed for
running this by hand from a terminal, so they're commented out below rather
than deleted — search for "PROVISION (disabled)" to find and re-enable them
if you build an app around this later.

    pip install openpyxl beautifulsoup4 lxml requests mcp python-dotenv playwright
    python -m playwright install chromium

    cp .env.example .env        # paste in OPENAI_API_KEY
    python contact_extractor.py --input urls.xlsx --output out/contacts.xlsx --dry-run
    python contact_extractor.py --input urls.xlsx --output out/contacts.xlsx

Run `python contact_extractor.py --help` for every flag.
"""
from __future__ import annotations

import argparse
import asyncio
import json
import logging
import re
import shutil
import sys
import time
from concurrent.futures import ThreadPoolExecutor
from contextlib import AsyncExitStack
from dataclasses import dataclass, field, fields
from html import unescape
from pathlib import Path
from typing import Any, Protocol
from urllib.parse import parse_qsl, unquote, urldefrag, urlencode, urljoin, urlsplit, urlunsplit

import requests
from bs4 import BeautifulSoup, Tag
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    from dotenv import load_dotenv
    load_dotenv()  # no-op if there's no .env file
except ImportError:
    pass


# =============================================================================
# Config
# =============================================================================

class ConfigError(Exception):
    """Raised for anything that should stop the run before it starts."""


@dataclass
class ExtractionConfig:
    # --- required -----------------------------------------------------
    input_path: str = ""
    output_path: str | None = None  # None -> timestamped, next to input

    # --- model (OpenAI only) --------------------------------------------
    model: str = "gpt-4o-mini"
    api_key: str | None = None        # discouraged; prefer api_key_env / .env
    api_key_env: str = "OPENAI_API_KEY"

    # --- extraction behaviour -------------------------------------------
    mode: str = "hybrid"              # accepted for compatibility; shortlisted pages always go to the model
    llm_batch_size: int = 6
    llm_concurrency: int = 4
    retries: int = 3
    min_confidence: int = 25
    phone_region: str = "IN"

    # --- budget guardrails ------------------------------------------------
    # Per-URL cap: once a single input row's estimated model spend crosses
    # this, remaining model calls for that URL are skipped. Set to 0 to disable.
    max_cost_per_url: float | None = 0.15
    # Whole-run cap, checked between URLs — a second, coarser safety net.
    max_cost_usd: float | None = None

    # --- crawl policy -----------------------------------------------------
    browser_provider: str = "playwright"   # playwright | chrome-devtools | http
    mcp_command: str | None = None
    mcp_args: str | None = None
    # Shared across the whole run: if two input rows resolve to the same
    # domain, they draw from the same budget rather than each getting a
    # fresh 20 — see domain_page_counts in run().
    max_urls_per_domain: int = 20
    max_depth: int = 2
    per_domain_delay_ms: int = 700
    respect_robots: bool = True

    # PROVISION (disabled): reserved for running several sites at once.
    # The crawl uses one browser, so this isn't wired up to anything yet —
    # left as a field so turning it on later is one line.
    # site_concurrency: int = 1

    # --- run controls -----------------------------------------------------
    dry_run: bool = False
    overwrite: bool = False

    # PROVISION (disabled): resume-after-interruption support. See the
    # commented-out checkpoint functions further down for what this would
    # wire up to.
    # resume: str | None = None          # path to a checkpoint file to resume from

    # --- output / logging ---------------------------------------------
    log_file: str | None = None
    verbose: bool = False

    # PROVISION (disabled): NDJSON progress stream for an embedding
    # application to parse instead of human-readable log lines. See
    # emit_progress() further down.
    # progress_format: str = "text"      # text | ndjson

    # PROVISION (disabled): building this ExtractionConfig from a YAML file
    # via --config, so an embedding application can template a config
    # instead of passing a long argument list.
    #
    # @staticmethod
    # def from_yaml(path: str) -> dict[str, Any]:
    #     import yaml
    #     with open(path, encoding="utf-8") as fh:
    #         data = yaml.safe_load(fh) or {}
    #     valid = {f.name for f in fields(ExtractionConfig)}
    #     unknown = set(data) - valid
    #     if unknown:
    #         raise ConfigError(f"Unknown config key(s) in {path}: {', '.join(sorted(unknown))}")
    #     return data

    def resolve_api_key(self) -> str:
        if self.api_key:
            return self.api_key
        import os
        value = os.getenv(self.api_key_env)
        if not value:
            raise ConfigError(
                f"API key not found. Set the {self.api_key_env} environment variable "
                f"(or put it in a .env file), or pass --api-key for a one-off test."
            )
        return value

    def resolve_output_path(self) -> Path:
        if self.output_path:
            return Path(self.output_path)
        import datetime
        stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        stem = Path(self.input_path).stem or "contacts"
        return Path(self.input_path).parent / f"{stem}_contacts_{stamp}.xlsx"

    def validate(self) -> None:
        if not self.input_path:
            raise ConfigError("input_path is required.")
        if not Path(self.input_path).exists():
            raise ConfigError(f"Input file not found: {self.input_path}")
        if self.mode not in {"hybrid", "pure-llm"}:
            raise ConfigError(f"Unknown mode: {self.mode}")
        if self.max_urls_per_domain < 1:
            raise ConfigError("max_urls_per_domain must be at least 1.")
        if self.max_cost_per_url is not None and self.max_cost_per_url < 0:
            raise ConfigError("max_cost_per_url cannot be negative.")
        if not self.dry_run:
            self.resolve_api_key()  # raises ConfigError if missing


# =============================================================================
# Logging
# =============================================================================

logger = logging.getLogger("contact_extractor")


def setup_logging(verbose: bool = False, log_file: str | None = None) -> None:
    logger.handlers.clear()
    logger.setLevel(logging.DEBUG if verbose else logging.INFO)
    fmt = logging.Formatter("%(asctime)s  %(levelname)-7s  %(message)s", "%H:%M:%S")

    console = logging.StreamHandler(sys.stderr)
    console.setFormatter(fmt)
    logger.addHandler(console)

    if log_file:
        fh = logging.FileHandler(log_file, encoding="utf-8")
        fh.setFormatter(fmt)
        logger.addHandler(fh)


def emit_progress(event_type: str, **data: Any) -> None:
    """Just logs a human-readable line. The multi-file version also had an
    NDJSON-on-stdout mode and an on_progress() callback here for an
    embedding application to consume — see the ExtractionConfig.progress_format
    PROVISION comment above for what that looked like.
    """
    level = logging.INFO if event_type in {"site_start", "site_done", "site_skip", "run_done", "site_error"} else logging.DEBUG
    logger.log(level, "%s %s", event_type, {k: v for k, v in data.items() if k not in ("evidence",)})


# =============================================================================
# Shared data models
# =============================================================================

@dataclass
class PageSnapshot:
    url: str
    final_url: str
    title: str
    html: str
    role: str = "seed"
    error: str | None = None
    truncated: bool = False
    # Filled after fetch. contact_score is 0–100 from evidence on the page;
    # has_contact_details is true only when an email, phone, LinkedIn profile,
    # or schema.org person/contact point was actually found.
    contact_score: int = 0
    has_contact_details: bool = False


@dataclass
class Card:
    """A person-sized text block segmented out of a page, plus whatever
    identifiers (email/phone/LinkedIn) were found for it deterministically."""
    id: str
    page_url: str
    page_role: str
    text: str
    known_emails: list[str] = field(default_factory=list)
    known_phones: list[str] = field(default_factory=list)
    known_linkedin: str | None = None
    evidence: list[str] = field(default_factory=list)

    @property
    def is_anchored(self) -> bool:
        return bool(self.known_emails or self.known_phones or self.known_linkedin)


@dataclass
class ResolvedContact:
    """A contact fully resolved without a model call (schema.org)."""
    name: str | None
    title: str | None
    email: str | None
    phone: str | None
    linkedin: str | None
    company: str | None
    kind: str
    confidence: int
    evidence: list[str]
    source_url: str
    source_domain: str
    page_role: str = ""
    extraction_method: str = "schema.org"


@dataclass
class Contact:
    """The shape every contact ends up in, regardless of which path
    produced it (schema.org / rule+llm card / llm-only free text)."""
    name: str | None = None
    title: str | None = None
    email: str | None = None
    phone: str | None = None
    linkedin: str | None = None
    company: str | None = None
    kind: str = "person"                    # person | generic
    confidence: int = 0
    extraction_method: str = "rule+llm"      # schema.org | rule+llm | llm-only
    evidence: list[str] = field(default_factory=list)
    flagged_fields: list[str] = field(default_factory=list)
    source_url: str = ""       # the specific page this contact was found on
    source_domain: str = ""
    page_role: str = ""
    input_url: str = ""        # the original row from the input Excel

    def key(self) -> tuple:
        # A shared inbox (hello@, info@) is printed on every staff card.
        # Keying on the email alone would collapse those people into one row.
        if self.email and self.name:
            return ("email-name", self.email.lower(), self.name.lower())
        if self.email:
            return ("email", self.email.lower())
        if self.linkedin and self.name:
            return ("linkedin-name", self.linkedin.lower(), self.name.lower())
        if self.linkedin:
            return ("linkedin", self.linkedin.lower())
        if self.name:
            return ("name", self.name.lower(), self.source_domain)
        return ("phone", self.phone or "", self.source_domain)

    def merge(self, other: "Contact") -> None:
        for f in ("name", "title", "email", "phone", "linkedin", "company"):
            if not getattr(self, f) and getattr(other, f):
                setattr(self, f, getattr(other, f))
        if other.kind == "person":
            self.kind = "person"
        self.confidence = max(self.confidence, other.confidence)
        for e in other.evidence:
            if e not in self.evidence:
                self.evidence.append(e)
        for f in other.flagged_fields:
            if f not in self.flagged_fields:
                self.flagged_fields.append(f)


@dataclass
class PromptUnit:
    """One thing to ask the model about — either an anchored card (identify
    name/title only) or a free-text block (identify every field)."""
    id: str
    text: str
    fields_requested: tuple[str, ...]


@dataclass
class ContactGuess:
    card_id: str
    name: str | None = None
    title: str | None = None
    email: str | None = None
    phone: str | None = None
    linkedin: str | None = None


@dataclass
class ModelResponse:
    contacts: list[ContactGuess] = field(default_factory=list)
    prompt_tokens: int = 0
    completion_tokens: int = 0
    raw_error: str | None = None


class ModelClient(Protocol):
    model: str

    def extract(self, units: list[PromptUnit]) -> ModelResponse: ...
    def estimate_cost(self, prompt_tokens: int, completion_tokens: int) -> float: ...


@dataclass
class URLRow:
    url: str
    company_hint: str | None = None
    max_pages: int | None = None   # overrides this row's share of max_urls_per_domain
    notes: str | None = None


@dataclass
class ExtractionReport:
    output_path: str
    total_sites: int
    succeeded: int
    failed: int
    total_contacts: int
    total_pages: int
    est_cost_usd: float
    seconds: float
    per_site: list[dict] = field(default_factory=list)
    contacts: list[dict] = field(default_factory=list)


# =============================================================================
# Shared text/identifier utilities
# =============================================================================

def clean_text(value: str) -> str:
    return re.sub(r"\s+", " ", unescape(value or "")).strip(" \t\r\n|·•-–—:,")


def ensure_http_url(url: str) -> str:
    """Turn a bare host such as 7im.co.uk into https://7im.co.uk.

    urlsplit treats a value with no scheme as a path, so the hostname comes
    back empty and every such row shares one page budget.
    """
    url = (url or "").strip()
    if not url or url.startswith(("http://", "https://")):
        return url
    if url.startswith("//"):
        return "https:" + url
    return "https://" + url


def registrable_domain(url: str) -> str:
    return (urlsplit(ensure_http_url(url)).hostname or "").lower().removeprefix("www.")


_DROP_QUERY = {
    "utm_source", "utm_medium", "utm_campaign", "utm_content", "utm_term", "utm_id",
    "gclid", "fbclid", "mc_cid", "mc_eid", "request_locale", "_ga", "yclid",
}


def canonical_url(url: str) -> str:
    """One key for the same page.

    Trailing slashes and campaign parameters were opening the same contact
    page several times and using up the per-domain budget.
    """
    parts = urlsplit(ensure_http_url(url))
    kept = [
        (key, value) for key, value in parse_qsl(parts.query, keep_blank_values=False)
        if key.lower() not in _DROP_QUERY
    ]
    path = parts.path or "/"
    if path != "/":
        path = path.rstrip("/")
    return urlunsplit((parts.scheme.lower(), parts.netloc.lower(), path, urlencode(kept), ""))


GENERIC_LOCALS = {
    "info", "contact", "hello", "hi", "sales", "support", "help", "admin", "office",
    "enquiry", "enquiries", "inquiry", "inquiries", "careers", "jobs", "hr", "media",
    "press", "marketing", "billing", "accounts", "legal", "privacy", "webmaster",
    "noreply", "no-reply", "donotreply", "mail", "team", "general", "customercare",
    "ir", "investor", "investors", "reception", "frontdesk", "service", "bookings",
    "orders", "partners", "newsletter", "subscribe", "abuse", "dpo", "grievance",
}
JUNK_EMAIL_SUFFIX = re.compile(r"\.(png|jpe?g|gif|svg|webp|css|js|woff2?|ico)$", re.I)
JUNK_EMAIL_DOMAINS = re.compile(r"(sentry\.io|example\.(com|org)|domain\.com|yourdomain|email\.com|wixpress|placeholder)", re.I)
EMAIL_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,24}")
OBFUSCATED_RE = re.compile(
    r"([A-Za-z0-9._%+\-]+)\s*(?:\[\s*at\s*\]|\(\s*at\s*\)|\s+at\s+|&#64;|\{at\})\s*"
    r"([A-Za-z0-9.\-]+)\s*(?:\[\s*dot\s*\]|\(\s*dot\s*\)|\s+dot\s+|\{dot\})\s*([A-Za-z]{2,24})",
    re.I,
)
PHONE_TEXT_RE = re.compile(r"(?:\+|00)?\d[\d\s().\-]{7,20}\d")
LI_PERSON_RE = re.compile(r"https?://(?:[\w.]+\.)?linkedin\.com/in/[^\s\"'<>)]+", re.I)
LI_COMPANY_RE = re.compile(r"https?://(?:[\w.]+\.)?linkedin\.com/company/[^\s\"'<>)]+", re.I)
CARD_HINT = re.compile(
    r"(team|member|person|people|staff|profile|card|bio|employee|leader|management|"
    r"director|contact|vcard|author|agent|attorney|doctor|faculty)", re.I,
)
DIAL = {"IN": "91", "US": "1", "GB": "44", "AE": "971", "SG": "65"}
NSN = {"IN": 10, "US": 10, "GB": 10, "AE": 9, "SG": 8}


def normalise_email(raw: str) -> str | None:
    email = clean_text(unquote(raw)).lower().strip(".,;:()<>\"'").split("?")[0]
    if not re.fullmatch(r"[a-z0-9._%+\-]+@[a-z0-9.\-]+\.[a-z]{2,24}", email):
        return None
    if JUNK_EMAIL_SUFFIX.search(email) or JUNK_EMAIL_DOMAINS.search(email):
        return None
    # Form placeholders ("xxx@xxx.xxx") show up in validation copy, not as contacts.
    if re.fullmatch(r"(x+|name|email|user|yourname|test)@(x+|email|example|domain|test)\.[a-z]{2,24}", email):
        return None
    return email


def looks_like_phone(raw: str) -> bool:
    """True when a text fragment is formatted like a phone, not an id or coordinate.

    Bare digit runs, decimals (`18.125 15.2655`), and date-shaped values
    (`2601-02-02`) pass the loose digit regex and were being stored as phones.
    `tel:` links skip this check and go straight to normalise_phone.
    """
    text = clean_text(raw)
    if not text or re.search(r"[A-Za-z]", text) or "." in text:
        return False
    if re.search(r"(?:19|20)\d{2}[-/]\d{2}", text):
        return False
    groups = re.findall(r"\d+", text)
    if not groups:
        return False
    if any(len(group) > 6 for group in groups) and not text.startswith(("+", "00")):
        return False
    if len(groups) >= 4 and sum(len(group) <= 2 for group in groups) >= 3:
        return False
    digits = "".join(groups)
    if not 8 <= len(digits) <= 15:
        return False
    compact = re.sub(r"\s+", "", text)
    if re.fullmatch(r"\+?\d{7,15}", compact):
        return compact.startswith("+") or compact.startswith("00")
    return bool(re.search(r"[()\-\s]", text))


def normalise_phone(raw: str, region: str = "IN") -> str | None:
    text = clean_text(unquote(raw)).replace("tel:", "")
    is_plus = bool(re.match(r"^\s*(\+|00)", text))
    digits = re.sub(r"\D", "", text)
    if not (7 <= len(digits) <= 15) or re.fullmatch(r"(\d)\1+", digits):
        return None
    if not is_plus and re.fullmatch(r"(19|20)\d{6}", digits):
        return None  # looks like a date, not a phone number
    if is_plus:
        if digits.startswith("00"):
            digits = digits[2:]
        return "+" + digits
    cc, want = DIAL.get(region, DIAL["IN"]), NSN.get(region, NSN["IN"])
    if digits.startswith(cc) and len(digits) == len(cc) + want:
        return "+" + digits
    digits = digits.lstrip("0")
    if len(digits) == want:
        return "+" + cc + digits
    if region == "IN" and 10 <= len(digits) <= 11:
        return "+" + cc + digits[-10:]
    return "+" + cc + digits if len(digits) >= 8 else None


def decode_cf_email(hexstr: str) -> str | None:
    try:
        data = bytes.fromhex(hexstr)
        key = data[0]
        return "".join(chr(b ^ key) for b in data[1:])
    except ValueError:
        return None


# =============================================================================
# Browser — Python Playwright by default, MCP only for chrome-devtools, HTTP fallback
# =============================================================================

_USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36"
)
# Same snapshot as the MCP path, returned in one evaluate. Chunking was only
# needed because MCP tool replies are size-capped.
_CAPTURE_PAYLOAD_JS = """() => {
  const root = document.documentElement.cloneNode(true);
  root.querySelectorAll(
    'script:not([type="application/ld+json"]),style,noscript,svg,iframe,canvas,video,audio,path'
  ).forEach(n => n.remove());
  const html = '<html>' + root.innerHTML + '</html>';
  const max = %(max)d;
  return {
    html: html.slice(0, max),
    truncated: html.length > max,
    url: location.href,
    title: document.title || ''
  };
}"""
_CAPTURE_JS = """() => {
  const root = document.documentElement.cloneNode(true);
  root.querySelectorAll(
    'script:not([type="application/ld+json"]),style,noscript,svg,iframe,canvas,video,audio,path'
  ).forEach(n => n.remove());
  const html = '<html>' + root.innerHTML + '</html>';
  window.__ce_snap = html.slice(0, %(max)d);
  return JSON.stringify({
    len: window.__ce_snap.length,
    truncated: html.length > %(max)d,
    url: location.href,
    title: document.title || ''
  });
}"""
_CHUNK_JS = "() => window.__ce_snap.slice(%d, %d)"

_PROVIDER_COMMANDS: dict[str, tuple[str, list[str]]] = {
    "playwright": ("npx", ["-y", "@playwright/mcp@latest", "--headless", "--isolated"]),
    "chrome-devtools": ("npx", ["-y", "chrome-devtools-mcp@latest", "--headless", "--isolated"]),
}
SNAPSHOT_MAX = 400_000
SNAPSHOT_CHUNK = 48_000
PAGE_SETTLE_S = 1.2
_ROBOTS_CACHE: dict[str, list[str]] = {}


def _npx() -> str:
    if sys.platform.startswith("win"):
        return shutil.which("npx.cmd") or shutil.which("npx") or "npx.cmd"
    return shutil.which("npx") or "npx"


def _unwrap(result: Any) -> str:
    parts = [getattr(b, "text", None) for b in getattr(result, "content", []) or []]
    raw = "\n".join(p for p in parts if p).strip()
    raw = "\n".join(ln for ln in raw.split("\n") if not ln.strip().startswith("###")).strip()
    if len(raw) >= 2 and raw[0] == '"' and raw[-1] == '"':
        try:
            return json.loads(raw)
        except json.JSONDecodeError:
            return raw[1:-1]
    return raw


class Browser:
    def __init__(self, cfg: ExtractionConfig):
        self.cfg = cfg
        self.provider = cfg.browser_provider
        self._stack: AsyncExitStack | None = None
        self._session = None
        self._nav_tool: tuple[str, str] | None = None
        self._eval_tool: tuple[str, str] | None = None
        self._http: requests.Session | None = None
        self._pw = None
        self._browser = None
        self._context = None
        self._page = None
        self.mode = "http"

    async def __aenter__(self) -> "Browser":
        if self.provider == "http":
            self._start_http()
        elif self.provider == "playwright" and not self.cfg.mcp_command:
            try:
                await self._start_playwright()
            except Exception as exc:  # noqa: BLE001 - degrade, don't die
                logger.warning("Playwright failed to start (%s); falling back to HTTP", exc)
                await self._close_playwright()
                self._start_http()
        else:
            try:
                await self._start_mcp()
            except Exception as exc:  # noqa: BLE001 - degrade, don't die
                logger.warning("MCP browser failed to start (%s); falling back to HTTP", exc)
                await self._close_mcp()
                self._start_http()
        return self

    async def __aexit__(self, *_exc: Any) -> None:
        await self._close_playwright()
        await self._close_mcp()

    def _start_http(self) -> None:
        self.mode = "http"
        self._http = requests.Session()
        self._http.headers.update({
            "User-Agent": _USER_AGENT,
            "Accept-Language": "en",
        })

    async def _start_playwright(self) -> None:
        from playwright.async_api import async_playwright

        self._pw = await async_playwright().start()
        self._browser = await self._pw.chromium.launch(headless=True)
        self._context = await self._browser.new_context(
            user_agent=_USER_AGENT,
            locale="en-US",
            viewport={"width": 1366, "height": 900},
        )
        self._page = await self._context.new_page()
        self.mode = "playwright"
        logger.info("Playwright Chromium ready")

    async def _close_playwright(self) -> None:
        for obj in (self._context, self._browser):
            if obj is None:
                continue
            try:
                await obj.close()
            except Exception:  # noqa: BLE001
                pass
        if self._pw is not None:
            try:
                await self._pw.stop()
            except Exception:  # noqa: BLE001
                pass
        self._pw = self._browser = self._context = self._page = None

    async def _start_mcp(self) -> None:
        from mcp import ClientSession, StdioServerParameters
        from mcp.client.stdio import stdio_client

        if self.cfg.mcp_command:
            cmd, args = self.cfg.mcp_command, (self.cfg.mcp_args or "").split()
        else:
            cmd, args = _PROVIDER_COMMANDS.get(self.provider, _PROVIDER_COMMANDS["playwright"])
            if cmd == "npx":
                cmd = _npx()

        self._stack = AsyncExitStack()
        read, write = await self._stack.enter_async_context(
            stdio_client(StdioServerParameters(command=cmd, args=args))
        )
        self._session = await self._stack.enter_async_context(ClientSession(read, write))
        await asyncio.wait_for(self._session.initialize(), timeout=90)

        tools = (await self._session.list_tools()).tools
        self._nav_tool = self._pick(tools, ("navigate",), ("url",))
        self._eval_tool = self._pick(tools, ("evaluate", "eval_script"), ("function", "script", "expression", "code"))
        if not self._nav_tool or not self._eval_tool:
            names = ", ".join(t.name for t in tools)
            raise RuntimeError(f"MCP server exposes no navigate/evaluate pair (tools: {names})")
        self.mode = f"mcp:{self.provider}"
        logger.info("MCP browser ready via %s -> %s / %s", cmd, self._nav_tool[0], self._eval_tool[0])

    @staticmethod
    def _pick(tools, name_hints, arg_hints) -> tuple[str, str] | None:
        for tool in tools:
            if not any(h in tool.name.lower() for h in name_hints):
                continue
            schema = getattr(tool, "inputSchema", None) or getattr(tool, "input_schema", None) or {}
            props = (schema or {}).get("properties", {}) or {}
            for arg in arg_hints:
                if arg in props:
                    return tool.name, arg
            if props:
                return tool.name, next(iter(props))
        return None

    async def _close_mcp(self) -> None:
        if self._stack:
            try:
                await self._stack.aclose()
            except Exception:  # noqa: BLE001
                pass
            self._stack, self._session = None, None

    async def open_page(self, url: str, role: str = "seed") -> PageSnapshot:
        try:
            if self._page is not None:
                return await self._open_playwright(url, role)
            if self._session:
                return await self._open_mcp(url, role)
            return await asyncio.to_thread(self._open_http, url, role)
        except Exception as exc:  # noqa: BLE001
            return PageSnapshot(url, url, "", "", role, error=f"{type(exc).__name__}: {exc}")

    async def _open_playwright(self, url: str, role: str) -> PageSnapshot:
        await self._page.goto(url, wait_until="domcontentloaded", timeout=45_000)
        try:
            await self._page.wait_for_load_state("load", timeout=8_000)
        except Exception:  # noqa: BLE001 - keep whatever has rendered
            pass
        await self._page.wait_for_timeout(int(PAGE_SETTLE_S * 1000))
        meta = await self._page.evaluate(_CAPTURE_PAYLOAD_JS % {"max": SNAPSHOT_MAX})
        if not isinstance(meta, dict):
            meta = {}
        return PageSnapshot(
            url=url,
            final_url=meta.get("url") or self._page.url or url,
            title=meta.get("title") or "",
            html=meta.get("html") or "",
            role=role,
            truncated=bool(meta.get("truncated")),
        )

    async def _call(self, tool: tuple[str, str], value: str) -> Any:
        name, arg = tool
        return await asyncio.wait_for(self._session.call_tool(name, {arg: value}), timeout=90)

    async def _open_mcp(self, url: str, role: str) -> PageSnapshot:
        await self._call(self._nav_tool, url)
        await asyncio.sleep(PAGE_SETTLE_S)
        meta_raw = _unwrap(await self._call(self._eval_tool, _CAPTURE_JS % {"max": SNAPSHOT_MAX}))
        try:
            meta = json.loads(meta_raw)
        except json.JSONDecodeError:
            start = meta_raw.find("{")
            meta = json.loads(meta_raw[start:]) if start >= 0 else {"len": 0, "url": url, "title": ""}
        total = int(meta.get("len", 0))
        parts = []
        for off in range(0, total, SNAPSHOT_CHUNK):
            parts.append(_unwrap(await self._call(self._eval_tool, _CHUNK_JS % (off, min(off + SNAPSHOT_CHUNK, total)))))
        return PageSnapshot(
            url=url, final_url=meta.get("url") or url, title=meta.get("title") or "",
            html="".join(parts), role=role, truncated=bool(meta.get("truncated")),
        )

    def _open_http(self, url: str, role: str) -> PageSnapshot:
        resp = self._http.get(url, timeout=30, allow_redirects=True)
        resp.raise_for_status()
        if "html" not in resp.headers.get("content-type", "text/html"):
            raise ValueError("not an HTML document")
        html = resp.text[:SNAPSHOT_MAX]
        return PageSnapshot(
            url=url, final_url=resp.url, title="", html=html, role=role,
            truncated=len(resp.text) > SNAPSHOT_MAX,
        )

    async def robots_allows(self, url: str) -> bool:
        if not self.cfg.respect_robots:
            return True
        parts = urlsplit(url)
        base = f"{parts.scheme}://{parts.netloc}"
        rules = _ROBOTS_CACHE.get(base)
        if rules is None:
            rules = []
            try:
                body = await asyncio.to_thread(lambda: requests.get(f"{base}/robots.txt", timeout=10).text)
                active = False
                for line in body.splitlines():
                    line = line.split("#")[0].strip()
                    if not line or ":" not in line:
                        continue
                    fld, value = (p.strip() for p in line.split(":", 1))
                    if fld.lower() == "user-agent":
                        active = value == "*"
                    elif fld.lower() == "disallow" and active and value:
                        rules.append(value)
            except Exception:  # noqa: BLE001 - no robots.txt = no restriction
                rules = []
            _ROBOTS_CACHE[base] = rules
        path = parts.path or "/"
        return not any(path.startswith(rule) for rule in rules)


# =============================================================================
# Crawl — link prior + page contact evidence
# =============================================================================
#
# Scoring is two layers:
#   1. Link prior (before fetch). Path, query, anchor text, aria-label and
#      title are scored for intent. Team, leadership, and staff pages are
#      opened before contact-us pages: a contact page usually has an address,
#      not names and titles. About sits between them so a leadership URL that
#      is only linked from About is still reached. Blog, product, career, and
#      similar paths are penalised.
#   2. Page evidence (after fetch). assess_page_contacts() only credits a page
#      when it actually contains an email, phone, LinkedIn profile, or
#      schema.org Person/ContactPoint. Keyword URLs that come back empty are
#      not expanded, so the per-domain budget is spent on pages that can
#      yield contacts. That check only decides whether to follow more links.
#      It does not extract the contacts — shortlisted pages go to the model.

# People pages outrank contact-us. Generic contact / reach-us / locations stay
# under the 70 follow-up floor so they are a fallback from the seed, not the
# next hop from a page that already lists staff.
LINK_RULES: list[tuple[re.Pattern[str], int, str]] = [
    (re.compile(r"\b(?:our[-_ ]?)?(?:team|people|staff|crew|colleagues|experts?)(?![-_a-z])", re.I), 100, "team"),
    (re.compile(r"\b(?:equipe|equipo|mannschaft|unser[-_ ]team|meet[-_ ]the[-_ ]team)\b", re.I), 100, "team"),
    (re.compile(r"\b(?:leadership|(?<![-_a-z])management|executives?|founders?|board|directors?|attorneys?|advisors?|faculty|physicians?|(?<![-_a-z])agents?)(?![-_a-z])", re.I), 100, "team"),
    (re.compile(r"\b(?:directory|who(?:s|')[-_ ]who|(?<![-_a-z])bios?|(?<![-_a-z])profiles?|speakers?)(?![-_a-z])", re.I), 96, "team"),
    (re.compile(r"\b(?:media|press)[-_ ](?:contact|relations|enquir|inquir|kit|office)\b", re.I), 76, "contact"),
    (re.compile(r"\b(?:investor|shareholder|ir)[-_ ](?:contact|relations)\b", re.I), 72, "contact"),
    (re.compile(
        r"\b(?:contacts?|kontakt|contacto|contato|contatti|nous[-_ ]contacter|contactez)"
        r"(?:[-_/ ](?:us|info|details|me|page|form|centre|center))?(?![-_a-z0-9])",
        re.I,
    ), 66, "contact"),
    (re.compile(r"\b(?:reach|get)[-_ ](?:us|in[-_ ]?touch)\b", re.I), 64, "contact"),
    (re.compile(r"\b(?:imprint|impressum|legal[-_ ]?notice|mentions[-_ ]?legales)\b", re.I), 60, "contact"),
    (re.compile(r"\b(?:about(?:[-_/ ]us)?|a[-_ ]propos|ueber[-_ ]uns|uber[-_ ]uns|chi[-_ ]siamo|quienes[-_ ]somos|sobre[-_ ]nosotros)(?![-_a-z])", re.I), 58, "about"),
    (re.compile(r"\b(?:offices?|locations?|branches?|find[-_ ]us|where[-_ ]to[-_ ]find)(?![-_a-z])", re.I), 50, "contact"),
    (re.compile(r"\b(?:sales|support|enquir\w*|inquir\w*|customer[-_ ]?care|helpdesk)(?![-_a-z])", re.I), 48, "contact"),
    (re.compile(r"\b(?:who[-_ ]we[-_ ]are|our[-_ ]story)\b", re.I), 42, "about"),
]
# Lower opens first. Team/leadership/staff, then about (a bridge to those
# pages), then contact-us and other address-style pages.
_OPEN_PRIORITY = {"team": 0, "about": 1, "contact": 2, "seed": 3, "other": 4}
# "partners" and bare "company" / "address" used to score well and pulled in
# partner-program, product, and blog URLs. They are no longer positive rules.
NEGATIVE_LINK = re.compile(
    r"\b(blog|news|articles?|stories|insights|products?|solutions?|services?|shop|store|"
    r"careers?|jobs?|vacanc|privacy|cookies?|terms|sitemap|events?|webinars?|podcasts?|"
    r"team[-_ ]building|sports?|gallery|downloads?|newsletter|login|signin|sign[-_ ]?up|signup)\b",
    re.I,
)
SKIP_LINK = re.compile(
    r"\.(pdf|docx?|xlsx?|pptx?|zip|rar|jpe?g|png|gif|svg|webp|mp4|mp3|ics)$|"
    r"(^|/)(login|signin|sign-in|sign-up|signup|register|cart|checkout|account|search|feed|rss|tag|"
    r"category|wp-admin|wp-content|privacy|terms|cookie|disclaimer|blog/\d{4})(/|$)",
    re.I,
)
# Open a link from the seed only when the prior clears this. Weaker matches
# ("who we are") still qualify; product/blog paths usually fall under it.
MIN_LINK_SCORE = 40
ADDRESS_RE = re.compile(
    r"\b\d{1,5}\s+[A-Za-z0-9.'\-]+(?:\s+[A-Za-z0-9.'\-]+){0,4}\s+"
    r"(?:Street|St|Road|Rd|Avenue|Ave|Lane|Ln|Drive|Dr|Boulevard|Blvd|Suite|Floor|Way)\b",
    re.I,
)
PEOPLE_HINT = re.compile(
    r"\b(chief|ceo|cfo|cto|coo|president|director|vice president|managing partner|"
    r"founder|head of|partner)\b",
    re.I,
)


@dataclass
class PageContactSignals:
    """What assess_page_contacts actually found on one HTML document."""
    emails: list[str] = field(default_factory=list)
    phones: list[str] = field(default_factory=list)
    linkedins: list[str] = field(default_factory=list)
    schema_people: int = 0
    schema_contacts: int = 0
    people_hints: int = 0
    has_address: bool = False
    has_contact_form: bool = False
    score: int = 0

    @property
    def has_details(self) -> bool:
        return bool(self.emails or self.phones or self.linkedins or self.schema_people or self.schema_contacts)


def _text_has_contact(text: str) -> bool:
    if not text:
        return False
    if EMAIL_RE.search(text) or OBFUSCATED_RE.search(text) or "linkedin.com/in/" in text.lower():
        return True
    return any(looks_like_phone(match) and normalise_phone(match) for match in PHONE_TEXT_RE.findall(text))


def assess_page_contacts(html: str, region: str = "IN") -> PageContactSignals:
    """Score a fetched page by the contact details it actually contains.

    Keyword overlap with the URL is intentionally ignored here. A page scores
    only from mailto/tel links (including mailto addresses escaped inside a
    script payload), visible emails and phones, LinkedIn profile URLs,
    Cloudflare-obfuscated emails, and schema.org Person/ContactPoint.
    """
    if not html:
        return PageContactSignals()
    soup = BeautifulSoup(html, "html.parser")
    emails: list[str] = []
    phones: list[str] = []
    linkedins: list[str] = []
    schema_people = 0
    schema_contacts = 0

    def add_email(raw: str | None) -> None:
        e = normalise_email(raw or "")
        if e and e not in emails:
            emails.append(e)

    def add_phone(raw: str | None) -> None:
        p = normalise_phone(raw or "", region)
        if p and p not in phones:
            phones.append(p)

    def add_linkedin(raw: str | None) -> None:
        if raw and "linkedin.com/in/" in raw.lower():
            url = raw.split("?")[0]
            if url not in linkedins:
                linkedins.append(url)

    # SPAs often keep the real mailto inside an escaped script
    # (&#x27;mailto:name@company.com) rather than an <a href>.
    for match in re.finditer(
        r"mailto:([A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,24})",
        unescape(html),
        re.I,
    ):
        add_email(match.group(1))

    for a in soup.find_all("a", href=True):
        href = a["href"].strip()
        low = href.lower()
        if low.startswith("mailto:"):
            add_email(href[7:])
        elif low.startswith("tel:"):
            add_phone(href[4:])
        else:
            add_linkedin(href)

    for tag in soup.select("[data-cfemail]"):
        decoded = decode_cf_email(tag.get("data-cfemail", ""))
        add_email(decoded)

    for script in soup.find_all("script", attrs={"type": "application/ld+json"}):
        try:
            nodes: list[dict] = []
            _walk_jsonld(json.loads(script.string or "{}"), nodes)
        except (json.JSONDecodeError, TypeError):
            continue
        for node in nodes:
            types = node.get("@type") or ""
            types = {types.lower()} if isinstance(types, str) else {str(t).lower() for t in types if isinstance(t, str)}
            if node.get("email"):
                add_email(str(node.get("email")))
            if node.get("telephone"):
                add_phone(str(node.get("telephone")))
            same_as = node.get("sameAs") or []
            same_as = [same_as] if isinstance(same_as, str) else same_as
            for item in same_as:
                if isinstance(item, str):
                    add_linkedin(item)
            if "person" in types:
                schema_people += 1
            if "contactpoint" in types or ("organization" in types and (node.get("email") or node.get("telephone"))):
                schema_contacts += 1

    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    text = soup.get_text(" ")
    for match in EMAIL_RE.findall(text):
        add_email(match)
    for match in OBFUSCATED_RE.finditer(text):
        add_email(f"{match.group(1)}@{match.group(2)}.{match.group(3)}")
    for match in PHONE_TEXT_RE.findall(text):
        if looks_like_phone(match):
            add_phone(match)
    for match in LI_PERSON_RE.findall(text):
        add_linkedin(match)

    has_form = bool(
        soup.find("form")
        and (
            soup.find("input", attrs={"type": "email"})
            or soup.find("form", action=re.compile(r"contact", re.I))
        )
    )
    has_address = bool(ADDRESS_RE.search(text))
    people_hints = len(PEOPLE_HINT.findall(text))

    score = (
        min(40, 15 * len(emails))
        + min(25, 10 * len(phones))
        + min(20, 15 * len(linkedins))
        + (20 if schema_people else 0)
        + (15 if schema_contacts else 0)
        + (8 if has_address else 0)
        + (8 if has_form else 0)
        + (10 if people_hints >= 2 else 0)
    )
    return PageContactSignals(
        emails=emails, phones=phones, linkedins=linkedins,
        schema_people=schema_people, schema_contacts=schema_contacts,
        people_hints=people_hints, has_address=has_address, has_contact_form=has_form,
        score=min(score, 100),
    )


# A match of one of these words is rejected when another word is glued on
# ("contact-lenses", "team building"). "contact-us" is fine because that
# suffix is part of the contact rule and gets consumed before this check.
_PREFIX_WORDS = re.compile(r"^(contact|contacts|team|about|support|people|sales)$", re.I)
_ALLOWED_TAIL = re.compile(r"^[-_ ](?:us|info|details|me|page|form|centre|center)\b", re.I)


def _rule_hit(haystack: str) -> tuple[int, str]:
    best, role = 0, "other"
    for pattern, points, link_role in LINK_RULES:
        if points <= best:
            continue
        for match in pattern.finditer(haystack):
            tail = haystack[match.end():]
            last = re.split(r"[\s/_-]+", match.group(0).strip())[-1]
            # "contact-lenses" and "team building" are not contact pages.
            # "contact-us" / "contact-page" are. A slash is the next path
            # segment ("/team/jane-doe") and stays valid.
            if _PREFIX_WORDS.fullmatch(last) and re.match(r"[-_ ][A-Za-z]", tail) and not _ALLOWED_TAIL.match(tail):
                continue
            best, role = points, link_role
            break
    return best, role


def score_link(href: str, text: str) -> tuple[int, str]:
    """Prior for whether a link is worth opening. This does not see the destination page.

    The path is scored on its own. Anchor text can raise the score only when
    it is itself a contact/team label. A blog, product, or careers path does
    not inherit a high score from an anchor that merely contains the word
    "team" or "contact".
    """
    parts = urlsplit(href)
    path = parts.path or "/"
    query = unquote(parts.query).replace("&", " ").replace("=", " ")
    path_score, path_role = _rule_hit(f"{path} {query}")
    anchor_score, anchor_role = _rule_hit(text)
    path_negative = bool(NEGATIVE_LINK.search(path) or NEGATIVE_LINK.search(query))
    if path_negative or NEGATIVE_LINK.search(text) or len(text.split()) > 6:
        anchor_score = 0
    if path_negative and path_score < 80:
        path_score, path_role = 0, "other"

    evidence, evidence_role = 0, "other"
    if EMAIL_RE.search(text) or OBFUSCATED_RE.search(text):
        evidence, evidence_role = 92, "contact"
    if PHONE_TEXT_RE.search(text) and evidence < 70:
        evidence, evidence_role = 70, "contact"
    if "linkedin.com/in/" in href.lower() or "linkedin.com/in/" in text.lower():
        evidence, evidence_role = max(evidence, 88), "team"

    keyword = max(path_score, anchor_score)
    score = max(keyword, evidence)
    if keyword and evidence:
        score = min(100, keyword + evidence // 3)
    if evidence > keyword:
        role = evidence_role
    elif path_score >= anchor_score:
        role = path_role
    else:
        role = anchor_role

    score -= min(path.count("/") * 2, 12)
    return max(score, 0), role


def _open_order(item: tuple[int, str, str, str]) -> tuple[int, int]:
    """Team and leadership links before about, about before contact-us."""
    score, _url, role, _text = item
    return (_OPEN_PRIORITY.get(role, 9), -score)


def discover_links(page: PageSnapshot, root_domain: str, follow_offsite: bool = False) -> list[tuple[int, str, str, str]]:
    soup = BeautifulSoup(page.html, "html.parser")
    seen: dict[str, tuple[int, str, str, str]] = {}
    for a in soup.find_all("a", href=True):
        href = a["href"].strip()
        if not href or href.startswith(("mailto:", "tel:", "javascript:", "#")):
            continue
        absolute = urldefrag(urljoin(page.final_url, href))[0]
        if not absolute.startswith(("http://", "https://")) or SKIP_LINK.search(absolute):
            continue
        if not follow_offsite and registrable_domain(absolute) != root_domain:
            continue
        text = clean_text(" ".join(
            part for part in (
                a.get_text(" "),
                str(a.get("aria-label") or ""),
                str(a.get("title") or ""),
            ) if part
        ))[:160]
        score, role = score_link(absolute, text)
        if score < MIN_LINK_SCORE:
            continue
        current = seen.get(absolute)
        if not current or score > current[0]:
            seen[absolute] = (score, absolute, role, text)
    return sorted(seen.values(), key=_open_order)


def company_name(page: PageSnapshot, fallback: str) -> str:
    raw = page.title
    if not raw:
        soup = BeautifulSoup(page.html, "html.parser")
        raw = clean_text(soup.title.string) if soup.title and soup.title.string else ""
    if raw:
        parts = [
            clean_text(p) for p in re.split(r"\s*[|\-–—:•·»]\s*", raw)
            if p and not re.fullmatch(r"\s*(home|welcome|official\s+site|homepage)\s*", p, re.I)
        ]
        parts = [p for p in parts if len(p) > 2]
        if parts:
            chosen = (min(parts, key=len) if len(parts) > 1 else parts[0])[:80]
            if len(chosen.split()) <= 6:
                return chosen
    return fallback


def _stamp_contacts(page: PageSnapshot, signals: PageContactSignals) -> None:
    page.contact_score = signals.score
    page.has_contact_details = signals.has_details


def _expansion_floor(role: str, signals: PageContactSignals) -> int | None:
    """Minimum link prior required to follow outbound links from this page.

    A page that already has contact details (or is an explicit team/contact
    URL) only spends more budget on stronger contact/team links — footer
    "About" links stay queued out. An about page with no details may still
    bridge to a real contact or leadership URL, but nothing weaker.
    """
    if signals.has_details or role in {"team", "contact"}:
        return 70
    if role in {"about", "seed"}:
        return 75
    return None


async def navigate_site(browser: Browser, cfg: ExtractionConfig, seed_url: str) -> tuple[list[PageSnapshot], str, str | None]:
    """Returns (pages opened, detected company name, error-or-None)."""
    seed_url = ensure_http_url(seed_url)
    domain = registrable_domain(seed_url)

    if not await browser.robots_allows(seed_url):
        return [], domain, "Blocked by robots.txt"

    seed = await browser.open_page(seed_url, role="seed")
    if seed.error or not seed.html:
        return [], domain, seed.error or "Empty response"

    seed_signals = assess_page_contacts(seed.html, cfg.phone_region)
    _stamp_contacts(seed, seed_signals)
    logger.info(
        "opened %s role=seed contact_score=%d emails=%d phones=%d linkedin=%d",
        seed.final_url, seed_signals.score, len(seed_signals.emails),
        len(seed_signals.phones), len(seed_signals.linkedins),
    )

    company = company_name(seed, domain)
    pages = [seed]
    queue: list[tuple[int, str, str, str]] = []
    visited = {canonical_url(seed.final_url), canonical_url(seed_url)}
    queued: set[str] = set()

    def enqueue(item: tuple[int, str, str, str], floor: int) -> None:
        score, url, role, text = item
        url = canonical_url(url)
        if score < floor or url in visited or url in queued:
            return
        queued.add(url)
        queue.append((score, url, role, text))

    for link in discover_links(seed, domain):
        enqueue(link, MIN_LINK_SCORE)
    # Once two pages with real contact details are open, drop leftover
    # about/contact-us links and finish team, leadership, and staff URLs.
    # Generic contact-us scores under 70, so it is not opened after that.
    contact_pages = 1 if seed_signals.has_details else 0
    budget = cfg.max_urls_per_domain - 1

    while queue and budget > 0:
        if contact_pages >= 2:
            queue = [item for item in queue if item[2] == "team" or item[0] >= 70]
            queued = {canonical_url(item[1]) for item in queue}
            if not queue:
                break
        queue.sort(key=_open_order)
        score, url, role, _anchor = queue.pop(0)
        queued.discard(canonical_url(url))
        if canonical_url(url) in visited:
            continue
        visited.add(canonical_url(url))
        if not await browser.robots_allows(url):
            continue
        await asyncio.sleep(cfg.per_domain_delay_ms / 1000)
        sub = await browser.open_page(url, role=role)
        budget -= 1
        if sub.error or not sub.html:
            logger.debug("skip %s (%s)", url, sub.error)
            continue
        sub_signals = assess_page_contacts(sub.html, cfg.phone_region)
        _stamp_contacts(sub, sub_signals)
        logger.info(
            "opened %s role=%s link_score=%d contact_score=%d emails=%d phones=%d linkedin=%d",
            sub.final_url, role, score, sub_signals.score, len(sub_signals.emails),
            len(sub_signals.phones), len(sub_signals.linkedins),
        )
        pages.append(sub)
        visited.add(canonical_url(sub.final_url))
        if sub_signals.has_details:
            contact_pages += 1
        if cfg.max_depth > 1 and budget > 0:
            floor = _expansion_floor(role, sub_signals)
            if floor is None:
                continue
            for child in discover_links(sub, domain):
                if child[0] >= floor:
                    enqueue(child, floor)

    return pages, company, None


# =============================================================================
# Rules — schema.org and person cards. Not used to extract contacts.
#
# process_page used to call this pass and only asked the model when the rules
# could not cover the page. Shortlisted pages now go straight to the model.
# _walk_jsonld is still used by assess_page_contacts while deciding which
# links to follow.
# =============================================================================

def _walk_jsonld(node, out: list[dict]) -> None:
    if isinstance(node, list):
        for item in node:
            _walk_jsonld(item, out)
    elif isinstance(node, dict):
        out.append(node)
        for value in node.values():
            if isinstance(value, (dict, list)):
                _walk_jsonld(value, out)


def extract_schema_org(soup: BeautifulSoup, page: PageSnapshot, company: str | None, region: str) -> list[ResolvedContact]:
    out: list[ResolvedContact] = []
    nodes: list[dict] = []
    for script in soup.find_all("script", attrs={"type": "application/ld+json"}):
        try:
            _walk_jsonld(json.loads(script.string or "{}"), nodes)
        except (json.JSONDecodeError, TypeError):
            continue
    domain = registrable_domain(page.final_url)
    for node in nodes:
        types = node.get("@type") or ""
        types = {types.lower()} if isinstance(types, str) else {t.lower() for t in types if isinstance(t, str)}
        email = normalise_email(str(node.get("email", ""))) if node.get("email") else None
        phone = normalise_phone(str(node.get("telephone", "")), region) if node.get("telephone") else None
        same_as = node.get("sameAs") or []
        same_as = [same_as] if isinstance(same_as, str) else same_as
        linkedin = next((s for s in same_as if isinstance(s, str) and "linkedin.com/in/" in s), None)
        if "person" in types:
            name = clean_text(str(node.get("name", "")))
            if not (name or email):
                continue
            out.append(ResolvedContact(
                name=name or None, title=clean_text(str(node.get("jobTitle", ""))) or None,
                email=email, phone=phone, linkedin=linkedin, company=company, kind="person",
                confidence=90, evidence=["schema.org Person"],
                source_url=page.final_url, source_domain=domain, page_role=page.role,
            ))
        elif ("contactpoint" in types or "organization" in types) and (email or phone):
            out.append(ResolvedContact(
                name=None, title=clean_text(str(node.get("contactType") or node.get("name") or "")) or None,
                email=email, phone=phone, linkedin=None, company=company, kind="generic",
                confidence=70, evidence=["schema.org ContactPoint"],
                source_url=page.final_url, source_domain=domain, page_role=page.role,
            ))
    return out


def _ancestors(tag: Tag, limit: int = 6):
    node, seen = tag, 0
    while node is not None and seen < limit:
        node = node.parent
        seen += 1
        if isinstance(node, Tag):
            yield node


def _card_for(anchor: Tag) -> Tag | None:
    best = None
    for ancestor in _ancestors(anchor):
        if ancestor.name in {"body", "html", "[document]"}:
            break
        text = clean_text(ancestor.get_text(" "))
        attrs = " ".join(ancestor.get("class", []) + [ancestor.get("id", "")])
        if len(text) > 900:
            break
        if CARD_HINT.search(attrs) or ancestor.name in {"li", "article", "tr", "figure"}:
            return ancestor
        if 25 <= len(text) <= 500 and not best:
            best = ancestor
    return best


def _identifiers(card: Tag, region: str) -> tuple[list[str], list[str], str | None]:
    emails, phones, linkedin = [], [], None
    for a in card.find_all("a", href=True):
        href = a["href"].strip()
        low = href.lower()
        if low.startswith("mailto:"):
            e = normalise_email(href[7:])
            if e:
                emails.append(e)
        elif low.startswith("tel:"):
            p = normalise_phone(href[4:], region)
            if p:
                phones.append(p)
        elif "linkedin.com/in/" in low and not linkedin:
            linkedin = href.split("?")[0]
    for tag in card.select("[data-cfemail]"):
        decoded = decode_cf_email(tag.get("data-cfemail", ""))
        e = normalise_email(decoded) if decoded else None
        if e and e not in emails:
            emails.append(e)
    text = card.get_text(" ")
    for match in EMAIL_RE.findall(text):
        e = normalise_email(match)
        if e and e not in emails:
            emails.append(e)
    for match in OBFUSCATED_RE.finditer(text):
        e = normalise_email(f"{match.group(1)}@{match.group(2)}.{match.group(3)}")
        if e and e not in emails:
            emails.append(e)
    if clean_text(text).count(" ") < 100:  # only scan short blocks for bare phone digits
        for match in PHONE_TEXT_RE.findall(text):
            if not looks_like_phone(match):
                continue
            p = normalise_phone(match, region)
            if p and p not in phones:
                phones.append(p)
    if not linkedin:
        found = LI_PERSON_RE.search(str(card))
        linkedin = found.group(0) if found else None
    return emails, phones, linkedin


def _block_for_contact(tag: Tag) -> Tag | None:
    """Smallest ancestor that still contains the identifier plus nearby name/title text."""
    node: Tag | None = tag if isinstance(tag, Tag) else None
    while isinstance(node, Tag) and node.name not in {"body", "html", "[document]"}:
        text = clean_text(node.get_text(" "))
        if len(text) > 550:
            break
        if 40 <= len(text) <= 550:
            return node
        node = node.parent
    return None


def _contact_blocks(soup: BeautifulSoup) -> list[Tag]:
    """Person-sized blocks that contain an email, phone, or LinkedIn profile.

    Rules previously started only from mailto/tel anchors, a class-name hint,
    or a table row, so plain-text contacts never became cards.
    """
    seeds: list[Tag] = []
    for text_node in soup.find_all(string=_text_has_contact):
        parent = text_node.parent
        if isinstance(parent, Tag) and parent.name not in {"script", "style", "noscript"}:
            seeds.append(parent)
    for tag in soup.select("[data-cfemail]"):
        if isinstance(tag, Tag):
            seeds.append(tag)
    candidates: list[Tag] = []
    for seed in seeds:
        block = _block_for_contact(seed)
        if block is not None and not any(block is c for c in candidates):
            candidates.append(block)
    return [n for n in candidates if not any(n is not other and other in n.descendants for other in candidates)]


def find_cards(soup: BeautifulSoup, page: PageSnapshot, region: str) -> list[Card]:
    cards: list[Card] = []
    seen_nodes: list[Tag] = []

    anchors = [
        a for a in soup.find_all("a", href=True)
        if a["href"].lower().startswith(("mailto:", "tel:")) or "linkedin.com/in/" in a["href"].lower()
    ]
    anchors += soup.select("[data-cfemail]")
    for anchor in anchors:
        node = _card_for(anchor) if anchor.name == "a" else anchor.parent
        if node is not None and not any(node is c for c in seen_nodes):
            seen_nodes.append(node)

    for tag in soup.find_all(["li", "article", "div", "tr"], class_=CARD_HINT):
        text = clean_text(tag.get_text(" "))
        if 20 <= len(text) <= 500 and not any(tag is c for c in seen_nodes):
            seen_nodes.append(tag)

    for row in soup.find_all("tr"):  # classless directory tables
        if any(row is c for c in seen_nodes):
            continue
        cells = row.find_all(["td", "th"], recursive=False)
        text = clean_text(row.get_text(" "))
        if 2 <= len(cells) <= 8 and 15 <= len(text) <= 400:
            seen_nodes.append(row)

    for tag in _contact_blocks(soup):
        if any(tag is c or tag in c.descendants for c in seen_nodes):
            continue
        if any(c in tag.descendants for c in seen_nodes):
            continue
        seen_nodes.append(tag)

    for i, node in enumerate(seen_nodes):
        emails, phones, linkedin = _identifiers(node, region)
        text = clean_text(node.get_text("\n"))[:800]
        if not text:
            continue
        evidence = []
        if emails:
            evidence.append("email found in markup")
        if phones:
            evidence.append("phone found in markup")
        if linkedin:
            evidence.append("LinkedIn profile in markup")
        # A CSS class named "card" is not a person. Keep the block only when
        # it holds an identifier or a name-and-title line.
        if not (emails or phones or linkedin) and _person_from_text(text) is None:
            continue
        cards.append(Card(
            id=f"{page.role}-{i}", page_url=page.final_url, page_role=page.role,
            text=text, known_emails=emails, known_phones=phones, known_linkedin=linkedin,
            evidence=evidence,
        ))
    return cards


def _person_from_text(text: str) -> tuple[str, str] | None:
    """Name plus title from a short line such as 'Jane Fraser, Chief Executive Officer'."""
    text = clean_text(text)
    if not text or len(text) > 180:
        return None
    match = re.match(
        r"([A-Z][a-z]+(?:\s+(?!Chief\b|Chair\b|CEO\b|CFO\b|CTO\b|COO\b|President\b|Founder\b|Managing\b|Executive\b|Head\b|DBE\b|Dame\b|Sir\b)[A-Z][a-z]+){1,2})"
        r"\b[, ]+(.{6,90})$",
        text,
    )
    if not match:
        return None
    title = clean_text(match.group(2))
    if not re.match(
        r"(?:DBE\s+|Dame\s+|Sir\s+)*(?:Chair|Chief|CEO|CFO|CTO|COO|President|Founder|Managing|Executive|Head)\b",
        title,
        re.I,
    ):
        return None
    return match.group(1), title


def _inline_hidden_contacts(soup: BeautifulSoup) -> None:
    """Copy addresses that exist only in markup into the visible text.

    Staff pages often use a button labelled "Email Us" whose real address is
    only in href="mailto:...". get_text() drops the href, so the model never
    sees the person's email and the check against the source text drops it.
    """
    for tag in soup.select("[data-cfemail]"):
        decoded = decode_cf_email(str(tag.get("data-cfemail") or ""))
        if decoded and decoded.lower() not in tag.get_text(" ").lower():
            tag.append(f" ({decoded})")

    for anchor in soup.find_all("a", href=True):
        href = unescape(str(anchor.get("href") or "")).strip()
        low = href.lower()
        if low.startswith("mailto:"):
            revealed = href.split(":", 1)[1].split("?")[0].strip()
        elif low.startswith("tel:"):
            revealed = href.split(":", 1)[1].split("?")[0].strip()
        elif "linkedin.com/in/" in low:
            revealed = href.split("?")[0].strip()
        else:
            continue
        if revealed and revealed.lower() not in anchor.get_text(" ").lower():
            anchor.append(f" ({revealed})")


def _section_until_next_heading(heading: Tag, later: list[Tag]) -> str:
    """Text from this heading through the biography, stopping at the next person."""
    nxt = later[0] if later else None
    parts: list[str] = []
    for node in heading.next_elements:
        if nxt is not None and node is nxt:
            break
        if isinstance(node, str):
            parts.append(node)
    return clean_text(" ".join(parts))


def _heading_blocks(soup: BeautifulSoup) -> list[str]:
    """One snippet per staff heading.

    A card whose own column holds the name, title, and biography is used as-is.
    When the name and title are separate headings ("Stewart Blake" in an h3,
    "Chairman & Founder" in an h4), that column contains both, so the name-only
    piece is too short to keep. The section is then read forward, including
    the h4, until the next h2 or h3.
    """
    headings = soup.find_all(["h2", "h3"])
    if len(headings) < 2:
        return []
    blocks: list[str] = []
    used: list[Tag] = []
    for index, heading in enumerate(headings):
        chosen: Tag = heading
        for ancestor in _ancestors(heading, limit=12):
            if ancestor.name in {"body", "html", "[document]"}:
                break
            if any(other is not heading for other in ancestor.find_all(["h2", "h3", "h4"])):
                break
            chosen = ancestor
        if any(chosen is earlier or chosen in earlier.descendants for earlier in used):
            continue
        text = clean_text(chosen.get_text("\n"))
        if len(text) < 40:
            text = _section_until_next_heading(heading, headings[index + 1:])
        if len(text) < 40:
            continue
        used.append(chosen)
        blocks.append(text[:4000])
    return blocks


def full_page_text(soup: BeautifulSoup, max_chars: int = 6000) -> str:
    # Footer and header are kept: published emails and phones usually live there.
    # Nav is dropped so the model is not handed the menu.
    for tag in soup(["nav", "script", "style", "noscript"]):
        tag.decompose()
    _inline_hidden_contacts(soup)
    return clean_text(soup.get_text("\n"))[:max_chars]


def contact_excerpt(soup: BeautifulSoup, max_chars: int = 5000) -> str | None:
    """Footer, address, and the blocks that actually contain identifiers.

    Used when the rule pass found contact details it could not attach to a
    person card. The rest of the marketing page is left out of the prompt.
    """
    chunks: list[str] = []
    seen: set[str] = set()

    def add(text: str, limit: int = 1500) -> None:
        cleaned = clean_text(text)[:limit]
        if len(cleaned) < 25 or cleaned in seen:
            return
        seen.add(cleaned)
        chunks.append(cleaned)

    for tag in soup.find_all(["footer", "address"]):
        add(tag.get_text("\n"))
    for tag in _contact_blocks(soup):
        add(tag.get_text("\n"), limit=800)
    text = "\n".join(chunks)[:max_chars]
    return text if len(text) >= 40 else None


def _append_card_contact(resolved: list[ResolvedContact], card: Card, company: str | None, domain: str, page: PageSnapshot) -> None:
    """Keep an email, phone, or LinkedIn URL even when no name was parsed.

    Those identifiers used to wait for the model. A dry run, or a model
    reply with no name, dropped them.
    """
    if not card.is_anchored:
        return
    parsed = _person_from_text(card.text)
    name, title = parsed if parsed else (None, None)
    emails = list(card.known_emails)
    phones = list(card.known_phones)
    linkedin = card.known_linkedin
    evidence = ["identifier published on the page"]
    if name:
        evidence.append("name and title written together on the page")

    def add(email: str | None, phone: str | None, link: str | None, person_name: str | None) -> None:
        if email and any(item.email == email for item in resolved):
            return
        if link and any(item.linkedin == link for item in resolved):
            return
        if not email and not link and phone and any(item.phone == phone and not item.email for item in resolved):
            return
        if not any([person_name, email, phone, link]):
            return
        resolved.append(ResolvedContact(
            name=person_name, title=title if person_name else None,
            email=email, phone=phone, linkedin=link,
            company=company, kind="person" if person_name else "generic",
            confidence=72 if person_name else 60,
            evidence=list(evidence),
            source_url=page.final_url, source_domain=domain, page_role=page.role,
            extraction_method="rule",
        ))

    if len(emails) <= 1 and len(phones) <= 1:
        add(emails[0] if emails else None, phones[0] if phones else None, linkedin, name)
        return
    for email in emails:
        add(email, phones[0] if len(phones) == 1 else None, None, name if len(emails) == 1 else None)
    for phone in phones:
        if any(item.phone == phone for item in resolved):
            continue
        add(None, phone, None, None)
    if linkedin:
        add(None, None, linkedin, name if not emails else None)


def process_page(page: PageSnapshot, company: str | None, cfg: ExtractionConfig) -> tuple[list[ResolvedContact], list[Card], str | None]:
    """Return page text for the model. Rule-based contact extraction is not used.

    The crawl has already shortlisted this URL. When the page is a sequence of
    staff headings, each heading is its own snippet. Otherwise the whole page
    text is sent. Names, titles, emails, and phones are read by the model.
    """
    del company, cfg
    if not page.html:
        return [], [], None
    soup = BeautifulSoup(page.html, "html.parser")
    for tag in soup(["nav", "script", "style", "noscript"]):
        tag.decompose()
    _inline_hidden_contacts(soup)
    blocks = _heading_blocks(soup)
    if len(blocks) >= 2:
        cards = [
            Card(id=f"{page.role}-{i}", page_url=page.final_url, page_role=page.role, text=text)
            for i, text in enumerate(blocks)
        ]
        return [], cards, None
    text = clean_text(soup.get_text("\n"))[:16000]
    if len(text) < 40:
        return [], [], None
    return [], [], text


# =============================================================================
# Model client — OpenAI only, plain HTTP, structured output
# =============================================================================

SYSTEM_PROMPT = (
    "You extract employee or contact information from webpage text for a business "
    "directory tool. You will be given one or more snippets, each with an id. "
    "Set card_id to that snippet id, and use the same card_id for every person "
    "found in that snippet. Do not number the people. "
    "Only fill in a field if it is written explicitly in that snippet's text — "
    "never guess, complete, infer, or autocomplete a name, email address, phone number, "
    "or LinkedIn URL from partial information or from general knowledge. "
    "If a snippet does not describe an identifiable person, omit it from your output entirely. "
    "If a snippet lists more than one person, return one entry per person, and do not "
    "stop after the first group. Include people who have a name and a job title even "
    "when they have no email address and no phone number. "
    "A postal address, office location, or generic company inbox with no person's name "
    "is not a contact — do not invent a name or title for it. "
    "Return every field you are asked for; use null for anything not explicitly present."
)

CONTACT_RESPONSE_SCHEMA = {
    "type": "object",
    "properties": {
        "contacts": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "card_id": {"type": "string"},
                    "name": {"type": ["string", "null"]},
                    "title": {"type": ["string", "null"]},
                    "email": {"type": ["string", "null"]},
                    "phone": {"type": ["string", "null"]},
                    "linkedin": {"type": ["string", "null"]},
                },
                "required": ["card_id", "name", "title", "email", "phone", "linkedin"],
                "additionalProperties": False,
            },
        },
    },
    "required": ["contacts"],
    "additionalProperties": False,
}


def build_user_prompt(units: list[PromptUnit]) -> str:
    parts = []
    for u in units:
        wanted = ", ".join(u.fields_requested)
        parts.append(
            f"### SNIPPET {u.id}\n"
            f"card_id for every person in this snippet: {u.id}\n"
            f"Fields to fill in if present: {wanted}\n"
            f"{u.text}\n"
        )
    return "\n".join(parts)


# Per-million-token prices, USD. Rough, for the cost estimate in Run Summary —
# not billing-accurate; update as OpenAI's pricing page changes.
_PRICES = {
    "gpt-4o-mini": (0.15, 0.60),
    "gpt-4o": (2.50, 10.00),
    "gpt-4.1-mini": (0.40, 1.60),
    "gpt-4.1": (2.00, 8.00),
}


class OpenAIClient:
    def __init__(self, model: str, api_key: str, base_url: str = "https://api.openai.com/v1"):
        self.model = model
        self.api_key = api_key
        self.base_url = base_url.rstrip("/")

    def extract(self, units: list[PromptUnit]) -> ModelResponse:
        payload = {
            "model": self.model,
            "temperature": 0,
            "messages": [
                {"role": "system", "content": SYSTEM_PROMPT},
                {"role": "user", "content": build_user_prompt(units)},
            ],
            "response_format": {
                "type": "json_schema",
                "json_schema": {"name": "contact_extraction", "strict": True, "schema": CONTACT_RESPONSE_SCHEMA},
            },
        }
        resp = requests.post(
            f"{self.base_url}/chat/completions",
            headers={"Authorization": f"Bearer {self.api_key}", "Content-Type": "application/json"},
            json=payload, timeout=60,
        )
        if resp.status_code != 200:
            return ModelResponse(raw_error=f"HTTP {resp.status_code}: {resp.text[:400]}")
        data = resp.json()
        try:
            content = data["choices"][0]["message"]["content"]
            parsed = json.loads(content)
            contacts = [ContactGuess(**c) for c in parsed.get("contacts", [])]
        except (KeyError, IndexError, json.JSONDecodeError, TypeError) as exc:
            return ModelResponse(raw_error=f"Could not parse model response: {exc}")
        usage = data.get("usage", {})
        return ModelResponse(
            contacts=contacts,
            prompt_tokens=usage.get("prompt_tokens", 0),
            completion_tokens=usage.get("completion_tokens", 0),
        )

    def estimate_cost(self, prompt_tokens: int, completion_tokens: int) -> float:
        in_price, out_price = _PRICES.get(self.model, (0.50, 1.50))
        return (prompt_tokens / 1_000_000) * in_price + (completion_tokens / 1_000_000) * out_price


def get_model_client(model: str, api_key: str) -> ModelClient:
    return OpenAIClient(model=model, api_key=api_key)


# =============================================================================
# extract_llm — turns cards + free-text fallbacks into model calls, verifies
# every field the model returns against the source text before trusting it,
# and enforces the per-URL budget cap between batches.
# =============================================================================

@dataclass
class UnitContext:
    text: str
    known_emails: list[str]
    known_phones: list[str]
    known_linkedin: str | None
    is_anchored: bool
    source_url: str
    source_domain: str
    page_role: str
    company: str | None
    phone_region: str = "IN"


def build_units(
    pages_cards: list[tuple[str, str, str, list[Card], str | None]],
    company: str | None,
    phone_region: str = "IN",
) -> tuple[list[PromptUnit], dict[str, UnitContext]]:
    units: list[PromptUnit] = []
    ctx_map: dict[str, UnitContext] = {}
    counter = 0
    for page_url, page_role, _unused, cards, fallback_text in pages_cards:
        domain = registrable_domain(page_url)
        for card in cards:
            counter += 1
            uid = f"u{counter}"
            # One identifier can be paired with the name the model reads.
            # Several identifiers on one card have to be requested too, or
            # every person would inherit the first email on the block.
            single_anchor = card.is_anchored and len(card.known_emails) <= 1 and len(card.known_phones) <= 1
            fields_ = ("name", "title") if single_anchor else ("name", "title", "email", "phone", "linkedin")
            units.append(PromptUnit(id=uid, text=card.text, fields_requested=fields_))
            ctx_map[uid] = UnitContext(
                text=card.text, known_emails=card.known_emails, known_phones=card.known_phones,
                known_linkedin=card.known_linkedin, is_anchored=card.is_anchored,
                source_url=page_url, source_domain=domain, page_role=page_role, company=company,
                phone_region=phone_region,
            )
        if fallback_text:
            counter += 1
            uid = f"u{counter}"
            units.append(PromptUnit(id=uid, text=fallback_text, fields_requested=("name", "title", "email", "phone", "linkedin")))
            ctx_map[uid] = UnitContext(
                text=fallback_text, known_emails=[], known_phones=[], known_linkedin=None, is_anchored=False,
                source_url=page_url, source_domain=domain, page_role=page_role, company=company,
                phone_region=phone_region,
            )
    return units, ctx_map


def _normalized_emails_in(text: str) -> set[str]:
    found = {normalise_email(m) for m in EMAIL_RE.findall(text)}
    for m in OBFUSCATED_RE.finditer(text):
        found.add(normalise_email(f"{m.group(1)}@{m.group(2)}.{m.group(3)}"))
    return {e for e in found if e}


def email_in_text(candidate: str | None, text: str) -> bool:
    cand = normalise_email(candidate or "")
    return bool(cand) and cand in _normalized_emails_in(text)


def phone_in_text(candidate: str | None, text: str, region: str) -> bool:
    cand = normalise_phone(candidate or "", region)
    if not cand:
        return False
    cand_digits = cand.lstrip("+")[-8:]
    for match in PHONE_TEXT_RE.findall(text):
        found = normalise_phone(match, region)
        if found and found.lstrip("+")[-8:] == cand_digits:
            return True
    return False


def linkedin_in_text(candidate: str | None, text: str) -> bool:
    if not candidate:
        return False
    cand = candidate.split("?")[0].rstrip("/").lower()
    for match in LI_PERSON_RE.findall(text):
        if match.split("?")[0].rstrip("/").lower() == cand:
            return True
    return False


def _soft_present(value: str | None, text: str) -> bool:
    if not value:
        return False
    tokens = [t.lower() for t in value.split() if len(t) > 1]
    hay = text.lower()
    return bool(tokens) and sum(1 for t in tokens if t in hay) >= max(1, len(tokens) - 1)


def _email_near_name(text: str, name: str, emails: list[str]) -> str | None:
    token = next((part for part in name.split() if len(part) > 1), "")
    if not token:
        return None
    idx = text.lower().find(token.lower())
    if idx < 0:
        return None
    window = text[max(0, idx - 80): idx + 180].lower()
    for email in emails:
        if email.lower() in window:
            return email
    return None


def _build_contact(guess: ContactGuess, ctx: UnitContext) -> Contact | None:
    flagged: list[str] = []
    evidence = [f"{ctx.page_role} page"]
    name = (guess.name or "").strip() or None
    title = (guess.title or "").strip() or None

    if ctx.is_anchored:
        if len(ctx.known_emails) <= 1:
            email = ctx.known_emails[0] if ctx.known_emails else None
        else:
            guessed = normalise_email(guess.email or "")
            email = guessed if guessed in ctx.known_emails else _email_near_name(ctx.text, name or "", ctx.known_emails)
        if len(ctx.known_phones) <= 1:
            phone = ctx.known_phones[0] if ctx.known_phones else None
        else:
            guessed_phone = (guess.phone or "").strip()
            phone = guessed_phone if phone_in_text(guessed_phone, ctx.text, ctx.phone_region) else None
            if phone:
                phone = normalise_phone(phone, ctx.phone_region)
        linkedin = ctx.known_linkedin
        evidence.append("deterministic identifier (mailto/tel/LinkedIn)")
        if name and not _soft_present(name, ctx.text):
            flagged.append("name (model paraphrase, not verbatim)")
        if title and not _soft_present(title, ctx.text):
            flagged.append("title (model paraphrase, not verbatim)")
        method = "rule+llm"
        confidence = 70 + 15 * bool(email) + 5 * bool(phone) + 5 * bool(linkedin) + 5 * bool(name)
        confidence -= 5 * len(flagged)
    else:
        email = (guess.email or "").strip() or None
        phone = (guess.phone or "").strip() or None
        linkedin = (guess.linkedin or "").strip() or None
        if email and not email_in_text(email, ctx.text):
            flagged.append("email (not found in source text — dropped)")
            email = None
        if phone and not phone_in_text(phone, ctx.text, ctx.phone_region):
            flagged.append("phone (not found in source text — dropped)")
            phone = None
        if linkedin and not linkedin_in_text(linkedin, ctx.text):
            flagged.append("linkedin (not found in source text — dropped)")
            linkedin = None
        if name and not _soft_present(name, ctx.text):
            flagged.append("name (model paraphrase, not verbatim)")
        if title and not _soft_present(title, ctx.text):
            flagged.append("title (model paraphrase, not verbatim)")
        evidence.append("free-text extraction (no deterministic identifier on this block)")
        method = "llm-only"
        confidence = 20 + 25 * bool(email) + 10 * bool(phone) + 15 * bool(linkedin) + 10 * bool(name)
        confidence -= 5 * len([f for f in flagged if "dropped" in f])

    if not any([name, email, phone, linkedin]):
        return None

    return Contact(
        name=name, title=title, email=email, phone=phone, linkedin=linkedin,
        company=ctx.company, kind="person" if name else "generic",
        confidence=max(0, min(confidence, 99)), extraction_method=method,
        evidence=evidence, flagged_fields=flagged,
        source_url=ctx.source_url, source_domain=ctx.source_domain, page_role=ctx.page_role,
    )


def _context_for_guess(guess: ContactGuess, ctx_map: dict[str, UnitContext], batch: list[PromptUnit]) -> UnitContext | None:
    """Find the source snippet for one model row.

    The model sometimes numbers people ("1", "2") instead of repeating the
    snippet id. A batch with one snippet can only have come from that snippet.
    """
    ctx = ctx_map.get((guess.card_id or "").strip())
    if ctx is not None:
        return ctx
    if len(batch) == 1:
        return ctx_map.get(batch[0].id)
    # A staff page is several snippets. The model sometimes numbers people
    # instead of copying the snippet id. Keep the row when the name is written
    # in exactly one snippet of this batch.
    name = (guess.name or "").strip().lower()
    if len(name) > 3:
        matches = [unit for unit in batch if name in unit.text.lower() and unit.id in ctx_map]
        if len(matches) == 1:
            return ctx_map[matches[0].id]
    return None


def run_llm_extraction(
    cfg: ExtractionConfig,
    client: ModelClient,
    pages_cards: list[tuple[str, str, str, list[Card], str | None]],
    company: str | None,
    cost_cap: float | None = None,
) -> tuple[list[Contact], dict]:
    """cost_cap: once this URL's running estimated spend reaches it, no
    further batches are submitted. Checked between batches, not pre-emptively
    per call, since a call's actual cost isn't known until its response arrives.
    """
    units, ctx_map = build_units(pages_cards, company, cfg.phone_region)
    usage = {"prompt_tokens": 0, "completion_tokens": 0, "calls": 0, "est_cost": 0.0, "capped": False}
    if not units:
        return [], usage

    batches = [units[i:i + cfg.llm_batch_size] for i in range(0, len(units), cfg.llm_batch_size)]

    def call_batch(batch: list[PromptUnit]):
        resp = None
        for attempt in range(cfg.retries):
            resp = client.extract(batch)
            if resp.raw_error is None:
                return resp
            logger.warning("model call failed (attempt %d/%d): %s", attempt + 1, cfg.retries, resp.raw_error)
            time.sleep(min(2 ** attempt, 10))
        return resp

    contacts: list[Contact] = []
    running_cost = 0.0
    with ThreadPoolExecutor(max_workers=max(1, cfg.llm_concurrency)) as pool:
        i = 0
        while i < len(batches):
            if cost_cap and running_cost >= cost_cap:
                usage["capped"] = True
                skipped = len(batches) - i
                logger.warning(
                    "per-URL budget cap ($%.4f) reached after $%.4f spent; skipping %d remaining batch(es)",
                    cost_cap, running_cost, skipped,
                )
                break
            chunk = batches[i:i + cfg.llm_concurrency]
            for batch, resp in zip(chunk, pool.map(call_batch, chunk)):
                usage["calls"] += 1
                if resp is None or resp.raw_error:
                    logger.error("batch failed permanently: %s", resp.raw_error if resp else "no response")
                    continue
                usage["prompt_tokens"] += resp.prompt_tokens
                usage["completion_tokens"] += resp.completion_tokens
                running_cost += client.estimate_cost(resp.prompt_tokens, resp.completion_tokens)
                for guess in resp.contacts:
                    ctx = _context_for_guess(guess, ctx_map, batch)
                    if ctx is None:
                        continue
                    contact = _build_contact(guess, ctx)
                    if contact:
                        contacts.append(contact)
            i += len(chunk)

    usage["est_cost"] = running_cost
    return contacts, usage


# =============================================================================
# Merge / dedup
# =============================================================================

_NAME_NOISE = re.compile(
    r"\b(?:mr|mrs|ms|miss|dr|prof|sir|dame|lord|lady|jr|sr|ii|iii|fcca|aca|cpa|mba|phd|esq)\b",
    re.I,
)


def _name_tokens(name: str) -> list[str]:
    cleaned = _NAME_NOISE.sub(" ", name.lower())
    cleaned = re.sub(r"[^a-z\s'-]", " ", cleaned)
    return [part for part in re.split(r"[\s'-]+", cleaned) if len(part) >= 2]


def _email_local_pieces(email: str) -> list[str]:
    local = email.split("@", 1)[0].lower().split("+", 1)[0]
    return [part for part in re.split(r"[._\-]+", local) if part]


def email_matches_name(name: str | None, email: str | None) -> bool:
    """True when the mailbox looks like this person's own address.

    hello@, info@, and sales@ do not. sean.edwards@ does for Sean Edwards,
    including first name only, last name only, and an initial plus the surname.
    """
    if not name or not email or "@" not in email:
        return False
    pieces = _email_local_pieces(email)
    if not pieces or all(part in GENERIC_LOCALS for part in pieces):
        return False
    tokens = _name_tokens(name)
    if not tokens:
        return False
    squashed = "".join(pieces)

    def aligns(token: str, piece: str) -> bool:
        if piece == token:
            return True
        return len(token) >= 3 and len(piece) >= 3 and (piece.startswith(token) or token.startswith(piece))

    first, last = tokens[0], tokens[-1]
    first_hit = any(aligns(first, part) or (len(part) == 1 and part == first[0]) for part in pieces)
    last_hit = any(aligns(last, part) or (len(part) == 1 and part == last[0]) for part in pieces)
    if len(tokens) == 1:
        return first_hit
    if first_hit and last_hit:
        return True
    if squashed in {first + last, last + first, first[0] + last, last + first[0], first + last[0]}:
        return True
    if len(pieces) == 1 and (aligns(first, pieces[0]) or aligns(last, pieces[0])):
        return True
    return False


def drop_mismatched_emails(contacts: list[Contact]) -> list[Contact]:
    """Drop a named person's email when it is not their own address.

    The name and title stay. A row with no name keeps its email, because a
    company inbox on its own is still a contact.
    """
    for contact in contacts:
        if not contact.name or not contact.email:
            continue
        if email_matches_name(contact.name, contact.email):
            continue
        note = "email (does not match the name — removed)"
        if note not in contact.flagged_fields:
            contact.flagged_fields.append(note)
        contact.email = None
        contact.confidence = max(30, contact.confidence - 25) if contact.name else contact.confidence
    return contacts


def merge_contacts(contacts: list[Contact], min_confidence: int) -> list[Contact]:
    contacts = drop_mismatched_emails(contacts)
    merged: dict[tuple, Contact] = {}
    for c in contacts:
        if c.confidence < min_confidence:
            continue
        key = c.key()
        if key in merged:
            merged[key].merge(c)
        else:
            merged[key] = Contact(**{**c.__dict__, "evidence": list(c.evidence), "flagged_fields": list(c.flagged_fields)})

    named = [c for c in merged.values() if c.name and c.email]
    for key, c in list(merged.items()):
        if c.name or not c.email:
            continue
        # Fold a nameless inbox into a person only when exactly one person
        # uses that address. A company hello@ printed on every card stays
        # its own row instead of being glued onto the first name.
        twins = [n for n in named if n.email and n.email.lower() == c.email.lower()]
        if len(twins) == 1:
            twins[0].merge(c)
            merged.pop(key, None)

    # After a shared inbox is stripped, the same person can be left as a
    # name-only row beside the row that still has their real address.
    with_email = [c for c in merged.values() if c.name and c.email]
    for key, c in list(merged.items()):
        if not c.name or c.email:
            continue
        twins = [
            n for n in with_email
            if n.name and n.name.lower() == c.name.lower() and n.source_domain == c.source_domain
        ]
        if len(twins) == 1:
            # The name-only row is what remains after a shared inbox was
            # stripped. Don't copy that removal note onto the row that still
            # has this person's own address.
            c.flagged_fields = [f for f in c.flagged_fields if "does not match the name" not in f]
            twins[0].merge(c)
            merged.pop(key, None)

    return sorted(
        merged.values(),
        key=lambda c: (c.kind != "person", -c.confidence, (c.name or "zz").lower()),
    )


# =============================================================================
# Excel I/O
# =============================================================================

_URL_HEADERS = {"url", "website", "link", "site", "company url", "company website"}


def read_urls(path: str) -> list[URLRow]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        raise ConfigError(f"{path} has no rows.")

    header = [str(c).strip().lower() if c else "" for c in rows[0]]
    url_col = next((i for i, h in enumerate(header) if h in _URL_HEADERS), None)
    has_header = url_col is not None
    if url_col is None:
        first_cell = str(rows[0][0] or "")
        url_col = 0
        has_header = "http" not in first_cell.lower() and "." not in first_cell

    company_col = header.index("company") if has_header and "company" in header else None
    max_pages_col = header.index("max_pages") if has_header and "max_pages" in header else None
    notes_col = header.index("notes") if has_header and "notes" in header else None

    data_rows = rows[1:] if has_header else rows
    out: list[URLRow] = []
    for row in data_rows:
        if url_col >= len(row) or not row[url_col]:
            continue
        url = ensure_http_url(str(row[url_col]).strip())
        if not url:
            continue
        out.append(URLRow(
            url=url,
            company_hint=str(row[company_col]).strip() if company_col is not None and row[company_col] else None,
            max_pages=int(row[max_pages_col]) if max_pages_col is not None and row[max_pages_col] else None,
            notes=str(row[notes_col]).strip() if notes_col is not None and row[notes_col] else None,
        ))
    if not out:
        raise ConfigError(f"No URLs found in {path}.")
    return out


CONTACT_COLUMNS = [
    "input_url", "company", "name", "title", "email", "phone", "linkedin",
    "kind", "confidence", "extraction_method", "source_url", "page_role",
    "evidence", "flagged_fields",
]
SUMMARY_COLUMNS = [
    "input_url", "status", "pages_opened", "pages_with_contacts", "contacts_found", "error",
    "model", "provider", "est_tokens", "est_cost_usd",
]
_HEADER_FILL = PatternFill("solid", fgColor="1F5E4A")
_HEADER_FONT = Font(color="FFFFFF", bold=True)


def _cell_value(value):
    if isinstance(value, list):
        return "; ".join(str(v) for v in value)
    return value


def _write_sheet(ws, columns: list[str], rows: list[dict]) -> None:
    ws.append(columns)
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
    for row in rows:
        ws.append([_cell_value(row.get(c)) for c in columns])
    for i, col in enumerate(columns, start=1):
        width = len(col)
        for r in rows:
            width = max(width, min(len(str(_cell_value(r.get(col, "")) or "")), 60))
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 60)
    ws.freeze_panes = "A2"


def resolve_output_path(path: Path, overwrite: bool) -> Path:
    if overwrite or not path.exists():
        return path
    stem, suffix = path.stem, path.suffix
    n = 2
    while True:
        candidate = path.with_name(f"{stem}_{n}{suffix}")
        if not candidate.exists():
            return candidate
        n += 1


def write_report(path: Path, contacts: list[Contact], summary_rows: list[dict], overwrite: bool = False) -> Path:
    final_path = resolve_output_path(path, overwrite)
    final_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws_contacts = wb.active
    ws_contacts.title = "Contacts"
    contact_rows = [
        {
            "input_url": c.input_url, "company": c.company, "name": c.name, "title": c.title,
            "email": c.email, "phone": c.phone, "linkedin": c.linkedin, "kind": c.kind,
            "confidence": c.confidence, "extraction_method": c.extraction_method,
            "source_url": c.source_url, "page_role": c.page_role,
            "evidence": c.evidence, "flagged_fields": c.flagged_fields,
        }
        for c in contacts
    ]
    _write_sheet(ws_contacts, CONTACT_COLUMNS, contact_rows)

    ws_summary = wb.create_sheet("Run Summary")
    _write_sheet(ws_summary, SUMMARY_COLUMNS, summary_rows)

    wb.save(final_path)
    return final_path


def _summary_row(input_url, status, pages_opened, contacts_found, error, cfg, est_tokens=0, est_cost=0.0, pages_with_contacts=0) -> dict:
    return {
        "input_url": input_url, "status": status, "pages_opened": pages_opened,
        "pages_with_contacts": pages_with_contacts,
        "contacts_found": contacts_found, "error": error, "model": cfg.model,
        "provider": "openai", "est_tokens": est_tokens, "est_cost_usd": round(est_cost, 4),
    }


def _resolved_to_contact(r: ResolvedContact) -> Contact:
    return Contact(
        name=r.name, title=r.title, email=r.email, phone=r.phone, linkedin=r.linkedin,
        company=r.company, kind=r.kind, confidence=r.confidence, extraction_method=r.extraction_method,
        evidence=list(r.evidence), flagged_fields=[], source_url=r.source_url,
        source_domain=r.source_domain, page_role=r.page_role,
    )


# =============================================================================
# PROVISION (disabled): checkpoint / --resume support.
#
# In the multi-file version, every completed site was appended as one JSON
# line to a `<output>.checkpoint.jsonl` file next to the output workbook.
# On a fresh run with --resume pointed at that file, already-completed URLs
# were skipped and their contacts/summary rows carried into the new output —
# so an interrupted long batch didn't mean re-paying for model calls already
# made. Left here as a full, working reference in case you want to wire it
# back in; nothing below this block calls it.
# =============================================================================
#
# def default_checkpoint_path(output_path: Path) -> Path:
#     return output_path.with_suffix(output_path.suffix + ".checkpoint.jsonl")
#
#
# def append_site_result(checkpoint_path: Path, input_url: str, contacts: list[Contact], summary: dict) -> None:
#     from dataclasses import asdict
#     checkpoint_path.parent.mkdir(parents=True, exist_ok=True)
#     with open(checkpoint_path, "a", encoding="utf-8") as fh:
#         fh.write(json.dumps({
#             "input_url": input_url,
#             "contacts": [asdict(c) for c in contacts],
#             "summary": summary,
#         }) + "\n")
#
#
# def load_checkpoint(checkpoint_path: str) -> tuple[set[str], list[Contact], list[dict]]:
#     path = Path(checkpoint_path)
#     if not path.exists():
#         return set(), [], []
#     done: set[str] = set()
#     contacts: list[Contact] = []
#     summaries: list[dict] = []
#     with open(path, encoding="utf-8") as fh:
#         for line in fh:
#             line = line.strip()
#             if not line:
#                 continue
#             record = json.loads(line)
#             done.add(record["input_url"])
#             contacts.extend(Contact(**c) for c in record["contacts"])
#             summaries.append(record["summary"])
#     return done, contacts, summaries


# =============================================================================
# Main pipeline
# =============================================================================

async def run_async(cfg: ExtractionConfig) -> ExtractionReport:
    setup_logging(cfg.verbose, cfg.log_file)
    cfg.validate()

    url_rows = read_urls(cfg.input_path)
    output_path = cfg.resolve_output_path()

    # PROVISION (disabled): resume support would load already-completed URLs
    # and their prior results here.
    #
    # already_done: set[str] = set()
    # carried_contacts: list[Contact] = []
    # carried_summaries: list[dict] = []
    # checkpoint_path = Path(cfg.resume) if cfg.resume else default_checkpoint_path(output_path)
    # if cfg.resume:
    #     already_done, carried_contacts, carried_summaries = load_checkpoint(cfg.resume)
    #     emit_progress("resume", already_done=len(already_done))
    already_done: set[str] = set()
    carried_contacts: list[Contact] = []
    carried_summaries: list[dict] = []

    client: ModelClient | None = None
    if not cfg.dry_run:
        client = get_model_client(cfg.model, cfg.resolve_api_key())
    # Used for cost *estimates* only (dry-run, and the per-URL budget check) —
    # never makes a network call, so this is safe even with no API key set.
    pricer = OpenAIClient(model=cfg.model, api_key="unused-for-pricing-only")

    started = time.time()
    emit_progress("run_start", sites=len(url_rows), provider="openai", model=cfg.model, mode=cfg.mode, dry_run=cfg.dry_run)

    all_contacts: list[Contact] = list(carried_contacts)
    summaries: list[dict] = list(carried_summaries)
    running_cost = sum(s.get("est_cost_usd", 0) or 0 for s in carried_summaries)
    total_pages = 0
    # Shared across every row this run: two input URLs on the same domain
    # draw from one budget (max_urls_per_domain) rather than each getting a
    # fresh one.
    domain_page_counts: dict[str, int] = {}

    async with Browser(cfg) as browser:
        emit_progress("browser_ready", mode=browser.mode)

        for row in url_rows:
            if row.url in already_done:
                emit_progress("site_skip", site=row.url, reason="already in checkpoint")
                continue

            if cfg.max_cost_usd and running_cost >= cfg.max_cost_usd:
                emit_progress("run_halt", reason="max_cost_usd reached", running_cost=running_cost)
                break

            domain = registrable_domain(row.url)
            used = domain_page_counts.get(domain, 0)
            remaining = cfg.max_urls_per_domain - used
            if remaining <= 0:
                msg = f"Per-domain URL cap reached ({cfg.max_urls_per_domain} for '{domain}')"
                summary = _summary_row(row.url, "skipped", 0, 0, msg, cfg)
                summaries.append(summary)
                # PROVISION (disabled): append_site_result(checkpoint_path, row.url, [], summary)
                emit_progress("site_skip", site=row.url, reason=msg)
                continue

            emit_progress("site_start", site=row.url)
            row_cap = min(row.max_pages, remaining) if row.max_pages else remaining
            site_cfg = ExtractionConfig(**{**cfg.__dict__, "max_urls_per_domain": row_cap})

            try:
                pages, company, error = await navigate_site(browser, site_cfg, row.url)
            except Exception as exc:  # noqa: BLE001 - one bad site must not kill the batch
                pages, company, error = [], row.url, f"{type(exc).__name__}: {exc}"

            domain_page_counts[domain] = used + len(pages)

            if error:
                summary = _summary_row(row.url, "failed", 0, 0, error, cfg)
                summaries.append(summary)
                # PROVISION (disabled): append_site_result(checkpoint_path, row.url, [], summary)
                emit_progress("site_error", site=row.url, message=error)
                continue

            company = row.company_hint or company
            total_pages += len(pages)
            pages_with_contacts = sum(1 for page in pages if page.has_contact_details)

            resolved_all: list[Contact] = []
            pages_cards = []
            for page in pages:
                resolved, cards, fallback_text = process_page(page, company, site_cfg)
                resolved_all.extend(_resolved_to_contact(r) for r in resolved)
                pages_cards.append((page.final_url, page.role, "", cards, fallback_text))

            if cfg.dry_run:
                unit_count = 0
                char_total = 0
                for _url, _role, _unused, cards, fallback_text in pages_cards:
                    unit_count += len(cards) + (1 if fallback_text else 0)
                    char_total += sum(len(c.text) for c in cards) + len(fallback_text or "")
                est_prompt_tokens = char_total // 4
                est_completion_tokens = unit_count * 40
                est_cost = pricer.estimate_cost(est_prompt_tokens, est_completion_tokens)
                site_contacts = resolved_all
                usage = {
                    "prompt_tokens": est_prompt_tokens, "completion_tokens": est_completion_tokens,
                    "calls": -(-unit_count // cfg.llm_batch_size) if unit_count else 0,
                    "est_cost": est_cost, "capped": False,
                }
                over_budget = bool(cfg.max_cost_per_url) and est_cost > cfg.max_cost_per_url
                emit_progress("site_dry_run", site=row.url, units=unit_count, est_prompt_tokens=est_prompt_tokens,
                               est_cost_usd=round(est_cost, 4), would_exceed_per_url_cap=over_budget)
            else:
                llm_contacts, usage = await asyncio.to_thread(
                    run_llm_extraction, cfg, client, pages_cards, company, cfg.max_cost_per_url,
                )
                site_contacts = resolved_all + llm_contacts
                running_cost += usage["est_cost"]
                if usage.get("capped"):
                    emit_progress("site_budget_capped", site=row.url, cap=cfg.max_cost_per_url, spent=usage["est_cost"])

            merged = merge_contacts(site_contacts, cfg.min_confidence)
            for c in merged:
                c.input_url = row.url
            all_contacts.extend(merged)

            status = "budget_capped" if usage.get("capped") else "ok"
            summary = _summary_row(
                row.url, status, len(pages), len(merged), None, cfg,
                est_tokens=usage["prompt_tokens"] + usage["completion_tokens"], est_cost=usage["est_cost"],
                pages_with_contacts=pages_with_contacts,
            )
            summaries.append(summary)
            # PROVISION (disabled): append_site_result(checkpoint_path, row.url, merged, summary)
            emit_progress("site_done", site=row.url, company=company, contacts=len(merged), pages=len(pages))

    final_path = write_report(output_path, all_contacts, summaries, overwrite=cfg.overwrite)
    seconds = round(time.time() - started, 1)
    report = ExtractionReport(
        output_path=str(final_path),
        total_sites=len(url_rows),
        succeeded=sum(1 for s in summaries if s["status"] in ("ok", "budget_capped")),
        failed=sum(1 for s in summaries if s["status"] == "failed"),
        total_contacts=len(all_contacts),
        total_pages=total_pages,
        est_cost_usd=round(running_cost, 4),
        seconds=seconds,
        per_site=summaries,
        contacts=[
            {
                "name": c.name, "title": c.title, "email": c.email,
                "phone": c.phone, "linkedin": c.linkedin, "company": c.company,
                "source_url": c.source_url, "input_url": c.input_url,
                "confidence": c.confidence, "page_role": c.page_role,
                "extraction_method": c.extraction_method,
            }
            for c in all_contacts
        ],
    )
    emit_progress("run_done", **{k: v for k, v in report.__dict__.items() if k != "per_site"})
    return report


# =============================================================================
# CLI
# =============================================================================

def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        prog="contact_extractor.py",
        description="Extract contact details from company/directory URLs listed in an Excel file, "
                     "navigating with Playwright Chromium and structuring results with OpenAI.",
    )
    p.add_argument("--input", required=True, help="Path to the input .xlsx with a URL column.")
    p.add_argument("--output", help="Path for the output .xlsx (default: timestamped, next to the input file).", default=argparse.SUPPRESS)

    # PROVISION (disabled): --config would load an ExtractionConfig.from_yaml(...)
    # file here, letting CLI flags override individual values on top of it.
    # p.add_argument("--config", help="Optional YAML config file; CLI flags override its values.", default=argparse.SUPPRESS)

    p.add_argument("--model", help="OpenAI model name. Default: gpt-4o-mini.", default=argparse.SUPPRESS)
    p.add_argument("--api-key", help="API key (plaintext — fine for a one-off test, not recommended otherwise).", default=argparse.SUPPRESS)
    p.add_argument("--api-key-env", help="Env var to read the OpenAI API key from. Default: OPENAI_API_KEY.", default=argparse.SUPPRESS)

    p.add_argument("--mode", choices=["hybrid", "pure-llm"],
                    help="Accepted for compatibility. Shortlisted pages are always sent "
                         "to the model; the rule-based contact pass is not used. Default: hybrid.",
                    default=argparse.SUPPRESS)
    p.add_argument("--llm-batch-size", type=int, help="Cards/snippets per model call. Default: 6.", default=argparse.SUPPRESS)
    p.add_argument("--llm-concurrency", type=int, help="Concurrent model calls. Default: 4.", default=argparse.SUPPRESS)
    p.add_argument("--retries", type=int, help="Retries per model call on transient failure. Default: 3.", default=argparse.SUPPRESS)
    p.add_argument("--min-confidence", type=int, help="Drop contacts scoring below this. Default: 25.", default=argparse.SUPPRESS)
    p.add_argument("--phone-region", help="Default region for bare phone numbers, e.g. IN, US. Default: IN.", default=argparse.SUPPRESS)

    p.add_argument("--browser-provider", choices=["playwright", "chrome-devtools", "http"],
                    help="Browser backend. 'playwright' launches Chromium through the Python "
                         "package. 'chrome-devtools' starts an MCP server via npx. 'http' fetches "
                         "HTML with no JavaScript. Default: playwright.",
                    default=argparse.SUPPRESS)
    p.add_argument("--mcp-command", help="Override the MCP server command.", default=argparse.SUPPRESS)
    p.add_argument("--mcp-args", help="Override the MCP server args (space-separated).", default=argparse.SUPPRESS)
    p.add_argument("--max-urls-per-domain", type=int, dest="max_urls_per_domain",
                    help="Max URLs navigated per domain, shared across every input row on that "
                         "domain in this run. Default: 20.", default=argparse.SUPPRESS)
    p.add_argument("--per-domain-delay-ms", type=int, help="Delay between page opens on the same site. Default: 700.", default=argparse.SUPPRESS)
    p.add_argument("--no-robots", dest="respect_robots", action="store_false", help="Ignore robots.txt (not recommended).", default=argparse.SUPPRESS)

    p.add_argument("--dry-run", action="store_true", help="Crawl and shortlist only; estimate model cost, spend nothing.", default=argparse.SUPPRESS)

    # PROVISION (disabled): --resume would point at a checkpoint .jsonl file.
    # p.add_argument("--resume", help="Path to a checkpoint .jsonl file to resume from.", default=argparse.SUPPRESS)

    p.add_argument("--max-cost-per-url", type=float,
                    help="Stop spending on a single URL once its estimated model cost crosses this. "
                         "Default: 0.15. Set to 0 to disable.", default=argparse.SUPPRESS)
    p.add_argument("--max-cost-usd", type=float, help="Stop the whole run once total estimated spend crosses this.", default=argparse.SUPPRESS)
    p.add_argument("--overwrite", action="store_true", help="Overwrite the output file instead of adding _2, _3, ...", default=argparse.SUPPRESS)

    # PROVISION (disabled): --progress-format ndjson would print one JSON
    # event per line to stdout instead of/alongside the human log lines, for
    # an embedding application to parse.
    # p.add_argument("--progress-format", choices=["text", "ndjson"], help="ndjson prints one JSON event per line to stdout.", default=argparse.SUPPRESS)

    p.add_argument("--log-file", help="Also write logs to this file.", default=argparse.SUPPRESS)
    p.add_argument("--verbose", action="store_true", help="Per-page debug logging.", default=argparse.SUPPRESS)
    return p


def _config_from_args(args: argparse.Namespace) -> ExtractionConfig:
    overrides = {}

    # PROVISION (disabled): --config YAML loading would merge in here first,
    # so CLI flags below still take precedence over the file.
    # config_path = getattr(args, "config", None)
    # if config_path:
    #     overrides.update(ExtractionConfig.from_yaml(config_path))

    # With default=argparse.SUPPRESS on every optional flag, vars(args) only
    # contains keys the user actually typed (plus the required --input) — so
    # an explicit `--no-robots` (value False) is correctly distinguished from
    # "flag not passed" rather than being silently dropped.
    valid_fields = {f.name for f in fields(ExtractionConfig)}
    cli_map = {"input": "input_path", "output": "output_path"}
    for key, value in vars(args).items():
        field_name = cli_map.get(key, key)
        if field_name in valid_fields:
            overrides[field_name] = value

    return ExtractionConfig(**overrides)


def main(argv: list[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)

    try:
        cfg = _config_from_args(args)
    except ConfigError as exc:
        print(f"Config error: {exc}", file=sys.stderr)
        return 2

    try:
        report = asyncio.run(run_async(cfg))
    except ConfigError as exc:
        print(f"Config error: {exc}", file=sys.stderr)
        return 2
    except KeyboardInterrupt:
        # PROVISION (disabled): with --resume wired back in, this message
        # would point at the checkpoint file next to the output path instead.
        print("\nInterrupted. Re-run the same command to start over — "
              "resume-from-interruption is disabled in this single-file build.", file=sys.stderr)
        return 1
    except Exception as exc:  # noqa: BLE001
        logger.exception("Run failed")
        print(f"Run failed: {exc}", file=sys.stderr)
        return 2

    print(f"\nDone in {report.seconds}s")
    print(f"  Sites: {report.succeeded} ok, {report.failed} failed, {report.total_sites} total")
    print(f"  Pages opened: {report.total_pages}")
    print(f"  Contacts: {report.total_contacts}")
    if not cfg.dry_run:
        print(f"  Estimated cost: ${report.est_cost_usd}")
    print(f"  Output: {report.output_path}")

    if report.total_sites == 0:
        return 3
    return 1 if report.failed else 0


if __name__ == "__main__":
    sys.exit(main())
