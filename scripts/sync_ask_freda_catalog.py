"""Rebuild the solution sheets of the Ask Freda catalog from the frontend.

The Playbooks pages and Ask Freda must offer the same solutions, so the
"Solutions Catalog", "Solution Sources" and "Solutions By Category" sheets of
data/Freda_Agents_and_Solutions_Catalog.xlsx (plus the Overview counts) and
backend/app/data/solutions_catalog.json are regenerated from frontend/src/data/any-site-reference/{datasets,vertical-datasets}.ts
via scripts/dump-solutions.mjs. Agent sheets are left untouched.

Runs automatically on backend startup; can also be run by hand:
    python scripts/sync_ask_freda_catalog.py
"""

from __future__ import annotations

import json
import logging
import shutil
import subprocess
import sys
import tempfile
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

logger = logging.getLogger("ask_freda_catalog_sync")

ROOT = Path(__file__).resolve().parents[1]
CATALOG_PATH = ROOT / "data" / "Freda_Agents_and_Solutions_Catalog.xlsx"
# Used by the backend's own catalog search (freda_catalog_service).
BACKEND_SOLUTIONS_JSON = ROOT / "backend" / "app" / "data" / "solutions_catalog.json"
DUMP_SCRIPT = ROOT / "scripts" / "dump-solutions.mjs"
SOURCES_DIR = ROOT / "frontend" / "src" / "data" / "any-site-reference"

SOLUTION_HEADERS = [
    "Solution ID", "Solution Name", "Category", "Tagline", "Description",
    "Rows / Records Available", "Coverage %", "Accuracy %", "Countries Covered",
    "Default Refresh Cadence", "Refresh Options", "# Sources", "Source Names",
    "# Output Attributes", "Output Attributes (Names)", "# Input Attributes", "Workflow ID",
]
SOURCE_HEADERS = [
    "Solution ID", "Solution Name", "Solution Category", "Source Name", "Source URL",
    "Source Kind", "Attributes Contributed", "Region",
]


def load_solutions() -> dict:
    node = shutil.which("node")
    if not node:
        raise RuntimeError("node is not on PATH; cannot read the frontend solution definitions")
    proc = subprocess.run(
        [node, "--no-warnings", str(DUMP_SCRIPT), str(SOURCES_DIR)],
        capture_output=True, text=True, encoding="utf-8", timeout=60,
    )
    if proc.returncode != 0:
        raise RuntimeError(f"dump-solutions.mjs failed: {proc.stderr.strip()[:500]}")
    data = json.loads(proc.stdout)
    if not data.get("datasets"):
        raise RuntimeError("dump-solutions.mjs returned no solutions")
    return data


def _replace_sheet(wb, name: str, headers: list[str], rows: list[list]) -> None:
    idx = wb.sheetnames.index(name)
    wb.remove(wb[name])
    ws = wb.create_sheet(name, idx)
    ws.append(headers)
    for row in rows:
        ws.append(row)


def sync(catalog_path: Path = CATALOG_PATH) -> dict:
    data = load_solutions()
    datasets = data["datasets"]

    solution_rows = []
    source_rows = []
    for d in datasets:
        sources = d.get("sources") or []
        outputs = d.get("outputAttributes") or []
        solution_rows.append([
            d["id"], d["name"], d["category"], d.get("tagline", ""), d.get("description", ""),
            d.get("rowsAvailable", ""), d.get("coverage"), d.get("accuracy"), d.get("countriesCovered"),
            d.get("refreshDefault", ""), "; ".join(d.get("refreshOptions") or []), len(sources),
            "; ".join(s["name"] for s in sources), len(outputs),
            "; ".join(a["label"] for a in outputs), len(d.get("inputAttributes") or []), d.get("workflowId", ""),
        ])
        for s in sources:
            source_rows.append([
                d["id"], d["name"], d["category"], s["name"], s.get("url") or None,
                s.get("kind") or None, s.get("attributes"), s.get("region") or None,
            ])

    by_category = Counter(d["category"] for d in datasets)
    category_rows = [[c, n] for c, n in sorted(by_category.items(), key=lambda kv: (-kv[1], kv[0]))]
    result = {"solutions": len(datasets), "categories": len(by_category), "sources": len(source_rows)}

    json_changed = _write_backend_json(datasets)
    # Only rewrite the workbook when the solution content actually differs, so
    # restarts don't leave the deployed git checkout with modified files.
    if not _sheets_differ(catalog_path, {
        "Solutions Catalog": [SOLUTION_HEADERS, *solution_rows],
        "Solution Sources": [SOURCE_HEADERS, *source_rows],
        "Solutions By Category": [["Category", "Solution Count"], *category_rows],
    }):
        return {**result, "changed": json_changed}

    # Write to a temp copy first so a failure never leaves a half-written catalog.
    wb = load_workbook(catalog_path)
    _replace_sheet(wb, "Solutions Catalog", SOLUTION_HEADERS, solution_rows)
    _replace_sheet(wb, "Solution Sources", SOURCE_HEADERS, source_rows)
    _replace_sheet(wb, "Solutions By Category", ["Category", "Solution Count"], category_rows)

    overview = wb["Overview"]
    updates = {
        "Total solutions (data products)": len(datasets),
        "Total solution categories": len(by_category),
        "Total solution source rows": len(source_rows),
        "Generated": datetime.now().replace(microsecond=0).isoformat(),
    }
    for row in overview.iter_rows(min_col=1, max_col=2):
        label = row[0].value
        if label in updates:
            row[1].value = updates[label]

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False, dir=catalog_path.parent) as tmp:
        tmp_path = Path(tmp.name)
    try:
        wb.save(tmp_path)
        tmp_path.replace(catalog_path)
    finally:
        tmp_path.unlink(missing_ok=True)

    return {**result, "changed": True}


def _norm(value):
    if value is None or value == "":
        return None
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def _sheets_differ(catalog_path: Path, expected: dict[str, list[list]]) -> bool:
    wb = load_workbook(catalog_path, read_only=True, data_only=True)
    try:
        for name, rows in expected.items():
            if name not in wb.sheetnames:
                return True
            current = [[_norm(v) for v in row] for row in wb[name].iter_rows(values_only=True)]
            current = [row for row in current if any(v is not None for v in row)]
            wanted = [[_norm(v) for v in row] for row in rows]
            if current != wanted:
                return True
        return False
    finally:
        wb.close()


def _write_backend_json(datasets: list[dict]) -> bool:
    """Write backend/app/data/solutions_catalog.json; True if it changed."""
    payload = {
        "solutions": [
            {
                "id": d["id"],
                "name": d["name"],
                "category": d["category"],
                "tagline": d.get("tagline", ""),
                "description": d.get("description", ""),
                "refreshDefault": d.get("refreshDefault", ""),
                "coverage": d.get("coverage"),
                "accuracy": d.get("accuracy"),
                "countriesCovered": d.get("countriesCovered"),
                "rowsAvailable": d.get("rowsAvailable", ""),
                "sources": [{"name": s["name"], "kind": s.get("kind") or None} for s in d.get("sources") or []],
                "outputFieldLabels": [a["label"] for a in d.get("outputAttributes") or []],
                "outputFieldCount": len(d.get("outputAttributes") or []),
            }
            for d in datasets
        ]
    }
    try:
        if json.loads(BACKEND_SOLUTIONS_JSON.read_text(encoding="utf-8")) == payload:
            return False
    except (OSError, ValueError):
        pass
    tmp = BACKEND_SOLUTIONS_JSON.with_suffix(".json.tmp")
    tmp.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    tmp.replace(BACKEND_SOLUTIONS_JSON)
    return True


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO)
    result = sync()
    state = "updated" if result["changed"] else "already up to date"
    print(f"Ask Freda catalog {state}: {result['solutions']} solutions, "
          f"{result['categories']} categories, {result['sources']} source rows -> {CATALOG_PATH}")
    sys.exit(0)
