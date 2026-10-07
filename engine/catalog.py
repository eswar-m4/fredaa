from __future__ import annotations

from functools import lru_cache
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
CATALOG_PATH = ROOT / "data" / "Freda_Agents_and_Solutions_Catalog.xlsx"


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _split(value: str) -> list[str]:
    return [part.strip() for part in value.replace("|", ";").split(";") if part.strip()]


def _hostname(url: str) -> str:
    raw = url.strip()
    if not raw or raw.startswith("{") or raw.startswith("("):
        return ""
    try:
        parsed = urlparse(raw if "://" in raw else f"https://{raw}")
        return (parsed.hostname or "").replace("www.", "").lower()
    except Exception:
        return ""


def _rows(ws) -> list[dict[str, str]]:
    values = list(ws.iter_rows(values_only=True))
    if not values:
        return []
    headers = [_cell(h) for h in values[0]]
    out: list[dict[str, str]] = []
    for row in values[1:]:
        record = {headers[i]: _cell(row[i] if i < len(row) else "") for i in range(len(headers)) if headers[i]}
        if any(record.values()):
            out.append(record)
    return out


@lru_cache(maxsize=1)
def load_catalog() -> dict:
    wb = load_workbook(CATALOG_PATH, data_only=True, read_only=True)
    overview_ws = wb["Overview"]
    overview: list[dict[str, str]] = []
    sheets: list[dict[str, str]] = []
    generated = ""
    legend = False
    for row in overview_ws.iter_rows(values_only=True):
        label, value = _cell(row[0] if row else ""), _cell(row[1] if row and len(row) > 1 else "")
        if not label:
            continue
        if label == "Generated":
            generated = value
            continue
        if label == "Sheet" and value == "Contents":
            legend = True
            continue
        if legend:
            sheets.append({"name": label, "contents": value})
        elif value:
            overview.append({"label": label, "value": value})

    agents = []
    for row in _rows(wb["Agents Catalog"]):
        if not row.get("Agent ID") or not row.get("Agent / Solution Name"):
            continue
        source = row.get("Source Link", "")
        agents.append(
            {
                "id": row["Agent ID"],
                "name": row["Agent / Solution Name"],
                "sourceUrl": source,
                "project": row.get("Project", ""),
                "type": row.get("Type", ""),
                "category": row.get("Category", ""),
                "industry": row.get("Industry", ""),
                "country": row.get("Country", ""),
                "dataType": row.get("Data Type Available", ""),
                "description": row.get("Description / Info", ""),
                "complexity": row.get("Complexity", ""),
                "datapoints": row.get("Datapoints", ""),
                "estimatedRecords": row.get("Estimated Records", ""),
                "runtimeType": row.get("Runtime Type", ""),
                "hostname": _hostname(source),
            }
        )

    solutions = []
    for row in _rows(wb["Solutions Catalog"]):
        if not row.get("Solution ID"):
            continue
        solutions.append(
            {
                "id": row["Solution ID"],
                "name": row.get("Solution Name", ""),
                "category": row.get("Category", ""),
                "tagline": row.get("Tagline", ""),
                "description": row.get("Description", ""),
                "records": row.get("Rows / Records Available", ""),
                "coverage": row.get("Coverage %", ""),
                "accuracy": row.get("Accuracy %", ""),
                "countriesCovered": row.get("Countries Covered", ""),
                "refreshCadence": row.get("Default Refresh Cadence", ""),
                "refreshOptions": _split(row.get("Refresh Options", "")),
                "sourceCount": row.get("# Sources", ""),
                "sourceNames": _split(row.get("Source Names", "")),
                "attributeCount": row.get("# Output Attributes", ""),
                "attributes": _split(row.get("Output Attributes (Names)", "")),
                "inputAttributeCount": row.get("# Input Attributes", ""),
            }
        )

    sources = []
    for row in _rows(wb["Solution Sources"]):
        if not row.get("Solution ID") or not row.get("Source Name"):
            continue
        sources.append(
            {
                "solutionId": row["Solution ID"],
                "solutionName": row.get("Solution Name", ""),
                "solutionCategory": row.get("Solution Category", ""),
                "name": row["Source Name"],
                "url": row.get("Source URL", ""),
                "kind": row.get("Source Kind", ""),
                "attributesContributed": row.get("Attributes Contributed", ""),
                "region": row.get("Region", ""),
            }
        )

    def facets(name: str) -> list[dict]:
        rows = _rows(wb[name])
        if not rows:
            return []
        keys = list(rows[0].keys())
        return [{"name": r.get(keys[0], ""), "count": int(float(r.get(keys[1], "0") or 0))} for r in rows if r.get(keys[0])]

    catalog = {
        "generated": generated,
        "overview": overview,
        "sheets": sheets,
        "agents": agents,
        "solutions": solutions,
        "sources": sources,
        "agentsByCategory": facets("Agents By Category"),
        "agentsByIndustry": facets("Agents By Industry"),
        "agentsByDataType": facets("Agents By Data Type"),
        "solutionsByCategory": facets("Solutions By Category"),
    }
    wb.close()
    return catalog
